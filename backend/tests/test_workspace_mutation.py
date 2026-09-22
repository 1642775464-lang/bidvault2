"""阶段六 C-3a：Material → Workspace 插入契约测试（不实现真实 OOXML mutation）。

契约见 阶段六C3a-插入契约设计.md：
- 插入=新 Workspace 正式版本（source=material_insert），与 ONLYOFFICE callback 无关
- base_workspace_version_id 乐观并发（409）；活跃编辑保护（423）；归属校验（400）
- workspace_mutations 审计表；服务不写 Source/Material；引擎未注入=501 且零副作用

注意：dispatch 同步注入（确定性）；替身引擎用 python-docx 追加段落（仅测试，非真实插入）。
"""

from __future__ import annotations

import hashlib
import io
import os
import time
from pathlib import Path

import pytest
from docx import Document as Docx
from fastapi.testclient import TestClient
from sqlalchemy.orm import sessionmaker

from app.db import get_session, make_engine
from app.main import app
from app.models import (
    AuditLog,
    Document,
    DocumentVersion,
    Material,
    OfficeDocument,
    OfficeDocumentVersion,
    WorkspaceDocument,
    WorkspaceMutation,
)


@pytest.fixture()
def client(tmp_path, monkeypatch):
    os.environ["ACCESS_PASSWORD"] = ""
    os.environ["SOURCE_MAINTENANCE_MODE"] = "0"
    os.environ["DATA_DIR"] = str(tmp_path / "data")
    os.environ["SCAN_INTERVAL_SECONDS"] = "3600"
    import app.config as cfg
    import app.db as appdb

    monkeypatch.setattr(cfg, "_SETTINGS", None)
    monkeypatch.setattr(appdb, "_Engine", None)
    monkeypatch.setattr(appdb, "_SessionLocal", None)
    engine = make_engine(cfg.get_settings().db_path)
    app.dependency_overrides[get_session] = _session_factory(engine)
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def _session_factory(engine):
    factory = sessionmaker(bind=engine, expire_on_commit=False)

    def _make():
        s = factory()
        try:
            yield s
        finally:
            s.close()

    return _make


def _db_session():
    """当前测试库的独立短会话（不碰 appdb 全局单例）。"""
    import app.config as cfg

    return sessionmaker(bind=make_engine(cfg.get_settings().db_path), expire_on_commit=False)()


def _docx_bytes(content: str) -> bytes:
    buf = io.BytesIO()
    d = Docx()
    for line in content.split("\n"):
        d.add_paragraph(line)
    d.save(buf)
    return buf.getvalue()


def _setup_word_material(client) -> tuple[str, str]:
    """上传 source → 拆解 → 建素材 → clone 出 Material OfficeDocument（v1, source=extract）。"""
    source = _docx_bytes("投标人须知前附表\n递交截止时间：2026-10-01 09:00\n")
    pid = client.post("/api/projects", json={"name": "C3a 素材项目"}).json()["id"]
    r = client.post(
        "/api/documents/upload",
        files=[("files", ("源文件.docx", source, "application/octet-stream"))],
        data={"kind": "tender", "project_id": pid, "is_secret": "false"},
    )
    doc_id = r.json()["created"][0]["id"]
    deadline = time.time() + 30
    while time.time() < deadline:
        docs = client.get("/api/documents").json()["items"]
        if next(x for x in docs if x["id"] == doc_id)["status"] == 2:
            break
        time.sleep(0.2)
    client.get(f"/api/documents/{doc_id}/editor-config")  # v1 建档
    m = client.post("/api/materials", json={"title": "C3a 素材", "cat": "投标格式", "type": "文本"}).json()
    r2 = client.post(
        f"/api/materials/{m['id']}/office-document",
        json={"source_document_id": doc_id, "child_range": [0, 1], "comment": "提取"},
    )
    assert r2.status_code == 200, r2.text
    return m["id"], r2.json()["office_document_id"]


def _setup_workspace(client) -> str:
    pid = client.post("/api/projects", json={"name": "C3a 工作稿项目"}).json()["id"]
    r = client.post(f"/api/projects/{pid}/workspace")
    assert r.status_code == 200, r.text
    return r.json()["office_document_id"]


def _version_id(oid: str, version: int) -> str:
    s = _db_session()
    try:
        v = (
            s.query(OfficeDocumentVersion)
            .filter(OfficeDocumentVersion.office_document_id == oid,
                    OfficeDocumentVersion.version == version)
            .one_or_none()
        )
    finally:
        s.close()
    assert v is not None, f"版本不存在：{oid[:8]} v{version}"
    return v.id


class _AppendEngine:
    """替身引擎：合法 DOCX 变更（追加素材标记段落）。真实 OOXML 包级插入属 C-3b。"""

    def mutate(self, plan, workspace_bytes: bytes, material_bytes: bytes) -> bytes:  # noqa: ARG002
        d = Docx(io.BytesIO(workspace_bytes))
        d.add_paragraph(f"【素材插入】{plan.material.title} v{plan.material_version.version}")
        out = io.BytesIO()
        d.save(out)
        return out.getvalue()


def _service(engine_double=None, probe=None):
    from app.services.office_document_svc import dispatch_post_save
    from app.services.workspace_mutation import WorkspaceMutationService

    def _sync_dispatch(oid: str, version_id: str) -> None:
        s = _db_session()
        try:
            od = s.get(OfficeDocument, oid)
            v = s.get(OfficeDocumentVersion, version_id)
            if od is not None and v is not None:
                dispatch_post_save(od, v)
        finally:
            s.close()

    return WorkspaceMutationService(
        session=_db_session(),
        active_edit_probe=probe or (lambda _oid: False),
        engine=engine_double,
        dispatch_fn=_sync_dispatch,
    )


def _req(ws_oid: str, base_vid: str, mid: str, mod_oid: str, mod_vid: str, **kw):
    from app.services.workspace_mutation import MaterialInsertRequest

    return MaterialInsertRequest(
        workspace_office_document_id=ws_oid,
        base_workspace_version_id=base_vid,
        material_id=mid,
        material_office_document_id=mod_oid,
        material_version_id=mod_vid,
        **kw,
    )


def _ep_body(_ws_oid: str, base_vid: str, mid: str, mod_oid: str, mat_vid: str, **kw):
    body = {"base_workspace_version_id": base_vid, "material_id": mid,
            "material_office_document_id": mod_oid, "material_version_id": mat_vid}
    body.update(kw)
    return body


def _fingerprint(s) -> tuple:
    """Source / Material 侧全量指纹（行数 + 关键列），用于"服务不写 Source/Material"断言。"""
    docs = [(d.id, d.status, d.current_version_id) for d in s.query(Document).all()]
    docver = [(v.id, v.file_hash, v.version) for v in s.query(DocumentVersion).all()]
    mats = [(m.id, m.title, m.text, m.rows_json, m.office_document_id) for m in s.query(Material).all()]
    return (tuple(docs), tuple(docver), tuple(mats))


# ---------------- 契约 1：完整链路（替身引擎） ----------------


def test_insert_happy_path_full_chain(client):
    mid, mod_oid = _setup_word_material(client)
    ws_oid = _setup_workspace(client)
    base_vid = _version_id(ws_oid, 1)
    mat_vid = _version_id(mod_oid, 1)

    s = _db_session()
    try:
        src_before = _fingerprint(s)
        matver_before = [(v.id, v.file_hash) for v in
                         s.query(OfficeDocumentVersion)
                         .filter(OfficeDocumentVersion.office_document_id == mod_oid).all()]
    finally:
        s.close()
    import app.routers.onlyoffice as oo

    cblog_before = dict(oo._office_cb_log)

    r = _service(engine_double=_AppendEngine()).insert_material(
        _req(ws_oid, base_vid, mid, mod_oid, mat_vid))
    assert r["ok"] is True and r["version"] == 2

    s = _db_session()
    try:
        ver_rows = (s.query(OfficeDocumentVersion)
                    .filter(OfficeDocumentVersion.office_document_id == ws_oid)
                    .order_by(OfficeDocumentVersion.version).all())
        current = s.get(OfficeDocument, ws_oid).current_version_id
        audit_rows = s.query(WorkspaceMutation).all()
        derived = (s.query(WorkspaceDocument)
                   .filter(WorkspaceDocument.office_document_id == ws_oid).one()
                   .derived_from_version_id)
        phase_log = s.query(AuditLog).filter(AuditLog.action == "workspace_material_insert").all()
    finally:
        s.close()

    assert [v.version for v in ver_rows] == [1, 2]
    assert ver_rows[1].source == "material_insert"
    assert "C3a 素材" in (ver_rows[1].comment or "")
    assert current == r["version_id"]
    assert len(audit_rows) == 1
    a = audit_rows[0]
    assert (a.workspace_office_document_id, a.workspace_base_version_id,
            a.workspace_new_version_id) == (ws_oid, base_vid, r["version_id"])
    assert (a.material_id, a.material_office_document_id, a.material_version_id) == (mid, mod_oid, mat_vid)
    assert a.insert_mode == "append_end"
    assert derived == r["version_id"]  # 同步 dispatch：派生已追平
    assert len(phase_log) == 1 and ws_oid in phase_log[0].detail

    # 不写 Source / Material；不伪造 ONLYOFFICE callback
    s = _db_session()
    try:
        assert _fingerprint(s) == src_before
        matver_after = [(v.id, v.file_hash) for v in
                        s.query(OfficeDocumentVersion)
                        .filter(OfficeDocumentVersion.office_document_id == mod_oid).all()]
    finally:
        s.close()
    assert matver_after == matver_before
    assert dict(oo._office_cb_log) == cblog_before


# ---------------- 契约 2：stale base → 409 ----------------


def test_stale_base_version_conflict(client):
    mid, mod_oid = _setup_word_material(client)
    ws_oid = _setup_workspace(client)
    v1 = _version_id(ws_oid, 1)
    mat_vid = _version_id(mod_oid, 1)

    svc = _service(engine_double=_AppendEngine())
    svc.insert_material(_req(ws_oid, v1, mid, mod_oid, mat_vid))  # v1 → v2

    with pytest.raises(Exception) as ei:
        svc.insert_material(_req(ws_oid, v1, mid, mod_oid, mat_vid))  # 仍基于 v1
    assert ei.value.code == "WORKSPACE_VERSION_CONFLICT" and ei.value.http_status == 409

    r = client.post(f"/api/office-documents/{ws_oid}/material-insert",
                    json=_ep_body(ws_oid, v1, mid, mod_oid, mat_vid))
    assert r.status_code == 409 and r.json()["detail"] == "WORKSPACE_VERSION_CONFLICT"


# ---------------- 契约 3：material 版本归属 ----------------


def test_material_version_mismatch(client):
    mid, mod_oid = _setup_word_material(client)
    ws_oid = _setup_workspace(client)
    base_vid = _version_id(ws_oid, 1)

    svc = _service(engine_double=_AppendEngine())
    # material_version_id 用了 workspace 的版本 → 不属于 material OD
    with pytest.raises(Exception) as ei:
        svc.insert_material(_req(ws_oid, base_vid, mid, mod_oid, base_vid))
    assert ei.value.code == "MATERIAL_VERSION_MISMATCH"
    # material_office_document_id 传了 workspace OD → 类型不符
    with pytest.raises(Exception) as ei:
        svc.insert_material(_req(ws_oid, base_vid, mid, ws_oid, base_vid))
    assert ei.value.code == "MATERIAL_VERSION_MISMATCH"

    r = client.post(f"/api/office-documents/{ws_oid}/material-insert",
                    json=_ep_body(ws_oid, base_vid, mid, mod_oid, base_vid))
    assert r.status_code == 400 and r.json()["detail"] == "MATERIAL_VERSION_MISMATCH"


# ---------------- 契约 4：workspace 版本归属 ----------------


def test_workspace_version_mismatch(client):
    mid, mod_oid = _setup_word_material(client)
    ws_oid = _setup_workspace(client)
    mat_vid = _version_id(mod_oid, 1)

    svc = _service(engine_double=_AppendEngine())
    # base 用了 material 的版本 → 不属于 workspace OD
    with pytest.raises(Exception) as ei:
        svc.insert_material(_req(ws_oid, mat_vid, mid, mod_oid, mat_vid))
    assert ei.value.code == "WORKSPACE_VERSION_MISMATCH"


# ---------------- 契约 5：活跃编辑保护（两路信号） ----------------


def test_active_edit_rejected(client):
    import app.config as cfg
    import app.routers.onlyoffice as oo

    mid, mod_oid = _setup_word_material(client)
    ws_oid = _setup_workspace(client)
    base_vid = _version_id(ws_oid, 1)
    mat_vid = _version_id(mod_oid, 1)

    # 信号 1：DS 会话态（内存，status=1 后未关闭）
    svc = _service(engine_double=_AppendEngine(), probe=lambda _oid: True)
    with pytest.raises(Exception) as ei:
        svc.insert_material(_req(ws_oid, base_vid, mid, mod_oid, mat_vid))
    assert ei.value.code == "WORKSPACE_EDITING" and ei.value.http_status == 423

    # 信号 2：工作快照脏（存在且 hash ≠ current）——status=6 未落版本的暂存
    snap = Path(cfg.get_settings().files_dir) / "working" / f"{ws_oid}.docx"
    snap.parent.mkdir(parents=True, exist_ok=True)
    snap.write_bytes(_docx_bytes("未落版本的暂存修改"))
    svc2 = _service(engine_double=_AppendEngine(), probe=lambda _oid: False)
    with pytest.raises(Exception) as ei:
        svc2.insert_material(_req(ws_oid, base_vid, mid, mod_oid, mat_vid))
    assert ei.value.code == "WORKSPACE_EDITING"

    # 快照与 current 一致（status=2 后的残留）→ 不算编辑中
    s = _db_session()
    try:
        cur = s.get(OfficeDocumentVersion, base_vid)
        cur_bytes = (Path(cfg.get_settings().files_dir) / cur.file_path).read_bytes()
    finally:
        s.close()
    snap.write_bytes(cur_bytes)
    assert hashlib.sha256(snap.read_bytes()).hexdigest() == cur.file_hash
    r = _service(engine_double=_AppendEngine(), probe=lambda _oid: False).insert_material(
        _req(ws_oid, base_vid, mid, mod_oid, mat_vid))
    assert r["version"] == 2

    # HTTP 层 423（v2 之后基于 v2、会话态=编辑中）
    oo._office_cb_log[ws_oid] = "编辑中"
    r2 = client.post(f"/api/office-documents/{ws_oid}/material-insert",
                     json=_ep_body(ws_oid, _version_id(ws_oid, 2), mid, mod_oid, mat_vid))
    assert r2.status_code == 423 and r2.json()["detail"] == "WORKSPACE_EDITING"


# ---------------- 契约 6+7：audit schema 与 source 枚举 ----------------


def test_audit_schema_and_source_enum(client):
    from app.models import Base
    from app.services.office_document_svc import OfficeDocumentService, OfficeDocumentTableAdapter

    # workspace_mutations 字段即契约字段；版本仍只在 office_document_versions（无新版本表）
    assert "workspace_mutations" in Base.metadata.tables
    cols = {c.name for c in Base.metadata.tables["workspace_mutations"].columns}
    assert cols == {"id", "workspace_office_document_id", "workspace_base_version_id",
                    "workspace_new_version_id", "material_id", "material_office_document_id",
                    "material_version_id", "insert_mode", "target_locator_json", "created_at"}
    assert not [n for n in Base.metadata.tables if "workspace_version" in n]

    mid, mod_oid = _setup_word_material(client)
    ws_oid = _setup_workspace(client)
    base_vid = _version_id(ws_oid, 1)
    mat_vid = _version_id(mod_oid, 1)

    # 引擎未注入：501（零副作用）；source 枚举契约另用 create_version 验证
    r = client.post(f"/api/office-documents/{ws_oid}/material-insert",
                    json=_ep_body(ws_oid, base_vid, mid, mod_oid, mat_vid))
    assert r.status_code == 501 and r.json()["detail"] == "MUTATION_ENGINE_NOT_IMPLEMENTED"

    s = _db_session()
    try:
        od = s.get(OfficeDocument, ws_oid)
        v = OfficeDocumentService(session=s, adapter=OfficeDocumentTableAdapter(s)).create_version(
            od, _docx_bytes("enum-check"), source="material_insert")
        s.expire(od)
        s.refresh(v)
    finally:
        s.close()
    assert v.source == "material_insert" and v.version == 2


# ---------------- 契约 8/9：insert_mode 校验、引擎未注入零副作用、404 ----------------


def test_invalid_insert_mode(client):
    mid, mod_oid = _setup_word_material(client)
    ws_oid = _setup_workspace(client)
    base_vid = _version_id(ws_oid, 1)
    mat_vid = _version_id(mod_oid, 1)
    r = client.post(f"/api/office-documents/{ws_oid}/material-insert",
                    json=_ep_body(ws_oid, base_vid, mid, mod_oid, mat_vid, insert_mode="after_heading"))
    assert r.status_code == 400 and r.json()["detail"] == "INVALID_INSERT_MODE"


def test_engine_missing_zero_side_effects(client):
    mid, mod_oid = _setup_word_material(client)
    ws_oid = _setup_workspace(client)
    base_vid = _version_id(ws_oid, 1)
    mat_vid = _version_id(mod_oid, 1)

    s = _db_session()
    try:
        ver_before = (s.query(OfficeDocumentVersion)
                      .filter(OfficeDocumentVersion.office_document_id == ws_oid).count())
        audit_before = s.query(WorkspaceMutation).count()
    finally:
        s.close()

    with pytest.raises(Exception) as ei:
        _service(engine_double=None).insert_material(_req(ws_oid, base_vid, mid, mod_oid, mat_vid))
    assert ei.value.code == "MUTATION_ENGINE_NOT_IMPLEMENTED" and ei.value.http_status == 501

    s = _db_session()
    try:
        ver_after = (s.query(OfficeDocumentVersion)
                     .filter(OfficeDocumentVersion.office_document_id == ws_oid).count())
        audit_after = s.query(WorkspaceMutation).count()
    finally:
        s.close()
    assert ver_after == ver_before and audit_after == audit_before


def test_workspace_od_not_found(client):
    r = client.post("/api/office-documents/nope/material-insert", json=_ep_body("x", "x", "x", "x", "x"))
    assert r.status_code == 404
