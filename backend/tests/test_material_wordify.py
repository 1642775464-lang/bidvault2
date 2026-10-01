"""素材溯源 v1.6：拆解素材盖章 child_range → 懒 Word 化（OOXML 保真 clone）→ 可插入工作稿。

- parser 行→body 子节点账本（line_children/image_children）
- 拆解管道为素材写 source_document_id/source_locator_json
- POST /materials/{id}/office-document 空体按溯源建档（幂等）
"""

from __future__ import annotations

import io
import json
import os
import time

import pytest
from docx import Document as Docx
from fastapi.testclient import TestClient
from sqlalchemy.orm import sessionmaker

from app.db import get_session, make_engine
from app.main import app


@pytest.fixture()
def client(tmp_path, monkeypatch):
    os.environ["ACCESS_PASSWORD"] = ""
    os.environ["SOURCE_MAINTENANCE_MODE"] = "0"
    os.environ["DATA_DIR"] = str(tmp_path / "data")
    os.environ["SCAN_INTERVAL_SECONDS"] = "3600"
    import app.config as cfg
    import app.db as appdb

    monkeypatch.setattr(cfg, "_SETTINGS", None)
    monkeypatch.setattr(appdb, "_Engine", None)
    monkeypatch.setattr(appdb, "_SessionLocal", None)
    engine = make_engine(cfg.get_settings().db_path)
    app.dependency_overrides[get_session] = _session_factory(engine)
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def _session_factory(engine):
    factory = sessionmaker(bind=engine, expire_on_commit=False)

    def _make():
        s = factory()
        try:
            yield s
        finally:
            s.close()

    return _make


def test_parser_line_children_accounting():
    from app.services.ingest.parser import parse_docx

    buf = io.BytesIO()
    d = Docx()
    d.add_paragraph("第一段")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "甲"
    t.cell(1, 1).text = "乙"
    d.add_paragraph("第二段")
    d.save(buf)
    parsed = parse_docx(buf.getvalue())
    # 段落行→自身子节点；表格每行→表格子节点（2 行同指子节点 1）
    assert len(parsed.line_children) == len(parsed.text.split("\n"))
    assert parsed.line_children[0] == 0
    assert parsed.line_children[1] == 1 and parsed.line_children[2] == 1
    assert parsed.line_children[3] == 2
    assert parsed.image_children == ()


def _upload_and_wait(client, name: str, body_docx: bytes, kind: str = "tender") -> str:
    r = client.post(
        "/api/documents/upload",
        files=[("files", (name, body_docx, "application/octet-stream"))],
        data={"kind": kind, "is_secret": "false"},
    )
    assert r.status_code == 200, r.text
    doc_id = r.json()["created"][0]["id"]
    deadline = time.time() + 40
    while time.time() < deadline:
        docs = client.get("/api/documents").json()["items"]
        if next(x for x in docs if x["id"] == doc_id)["status"] == 2:
            return doc_id
        time.sleep(0.3)
    raise AssertionError("拆解超时")


def test_pipeline_stamps_provenance_and_wordify(client):
    from app.services.ingest.parser import parse_docx

    buf = io.BytesIO()
    d = Docx()
    d.add_paragraph("投标人须知前附表")
    tb = d.add_table(rows=2, cols=2)
    tb.cell(0, 0).text = "条款号"
    tb.cell(0, 1).text = "内容"
    tb.cell(1, 0).text = "3.1"
    tb.cell(1, 1).text = "递交截止时间：2026-10-01 09:00"
    d.add_paragraph("评标办法前附表")
    tb2 = d.add_table(rows=1, cols=2)
    tb2.cell(0, 0).text = "评分项"
    tb2.cell(0, 1).text = "技术 60 分"
    d.save(buf)
    _upload_and_wait(client, "溯源样例.docx", buf.getvalue())

    mats = client.get("/api/materials").json()["items"]
    assert mats, "拆解未产出素材"
    stamped = [m for m in mats if m.get("source_document_id") and m.get("source_locator_json")]
    assert stamped, f"素材未盖溯源章: {[(m['title'], m.get('source_locator_json')) for m in mats]}"
    for m in stamped:
        loc = json.loads(m["source_locator_json"])
        assert isinstance(loc["child_range"], list) and len(loc["child_range"]) == 2
        assert loc["child_range"][0] <= loc["child_range"][1]

    # 懒 Word 化：空体按溯源建档（幂等）
    m0 = stamped[0]
    r1 = client.post(f"/api/materials/{m0['id']}/office-document", json={})
    assert r1.status_code == 200, r1.text
    oid = r1.json()["office_document_id"]
    assert oid and r1.json()["existed"] is False
    blocks = client.get(f"/api/office-documents/{oid}/reader").json()["blocks"]
    assert any(b["type"] == "table" for b in blocks), blocks
    # 幂等二连
    r2 = client.post(f"/api/materials/{m0['id']}/office-document", json={})
    assert r2.status_code == 200 and r2.json()["existed"] is True
    assert r2.json()["office_document_id"] == oid


def test_wordify_without_provenance_rejected(client):
    from app.models import Material

    import app.config as cfg
    import app.db as appdb
    appdb._Engine = None
    appdb._SessionLocal = None
    from sqlalchemy.orm import sessionmaker as _sm
    s = _sm(bind=make_engine(cfg.get_settings().db_path), expire_on_commit=False)()
    try:
        mat = Material(title="无源素材", type="文本", cat="其他", src="手工")
        s.add(mat)
        s.commit()
        mid = mat.id
    finally:
        s.close()
    _ = cfg
    r = client.post(f"/api/materials/{mid}/office-document", json={})
    assert r.status_code == 409


def test_clone_strips_comments():
    """批注属源文档评审层：clone 产物剥离批注锚点与参照 run（插入引擎拒绝清单不再命中）。"""
    import zipfile

    from app.services.ooxml_clone import clone_material

    buf = io.BytesIO()
    d = Docx()
    d.add_paragraph("投标人须知前附表")
    tb = d.add_table(rows=1, cols=2)
    tb.cell(0, 0).text = "A"
    tb.cell(0, 1).text = "B"
    d.add_paragraph("评标办法前附表")
    d.save(buf)
    src = buf.getvalue()
    # 注入批注锚点（模拟带批注的源文档：批注区间起 + 参照 run）
    zin = zipfile.ZipFile(io.BytesIO(src))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                x = data.decode("utf-8")
                x = x.replace(
                    "<w:sectPr",
                    '<w:commentRangeStart w:id="9"/>'
                    '<w:r><w:commentReference w:id="9"/></w:r><w:sectPr',
                    1,
                )
                data = x.encode("utf-8")
            zout.writestr(item, data)
    crafted = out.getvalue()

    cloned = clone_material(crafted, (0, 2), "带批注素材")
    z = zipfile.ZipFile(io.BytesIO(cloned))
    doc_xml = z.read("word/document.xml").decode("utf-8")
    assert "commentReference" not in doc_xml
    assert "commentRangeStart" not in doc_xml
    assert "评标办法前附表" in doc_xml
