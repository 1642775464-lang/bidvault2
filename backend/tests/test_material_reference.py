"""E-2C：引用闭环验收测试（C1 台账 / C2 插入复用 C-3 链 / C3 漂移 / C4 占用写回）。"""

from __future__ import annotations

import io
import os
import time
import zipfile

import pytest
from docx import Document as Docx
from fastapi.testclient import TestClient
from sqlalchemy.orm import sessionmaker


@pytest.fixture()
def client(tmp_path, monkeypatch):
    os.environ["ACCESS_PASSWORD"] = ""
    os.environ["DATA_DIR"] = str(tmp_path / "data")
    os.environ["SCAN_INTERVAL_SECONDS"] = "3600"
    import app.config as cfg
    import app.db as appdb

    monkeypatch.setattr(cfg, "_SETTINGS", None)
    monkeypatch.setattr(appdb, "_Engine", None)
    monkeypatch.setattr(appdb, "_SessionLocal", None)
    from app.db import get_session, make_engine
    from app.main import app

    monkeypatch.setattr(cfg, "_SETTINGS", None)
    engine = make_engine(cfg.get_settings().db_path)
    factory = sessionmaker(bind=engine, expire_on_commit=False)

    def _make():
        s = factory()
        try:
            yield s
        finally:
            s.close()

    app.dependency_overrides[get_session] = _make
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def _db():
    import app.config as cfg
    from app.db import make_engine

    return sessionmaker(bind=make_engine(cfg.get_settings().db_path), expire_on_commit=False)()


def _docx(lines: list[str]) -> bytes:
    buf = io.BytesIO()
    d = Docx()
    for ln in lines:
        if ln.startswith("#1 "):
            d.add_heading(ln[3:], level=1)
        else:
            d.add_paragraph(ln)
    d.save(buf)
    return buf.getvalue()


def _setup_ws_and_project(client, tmp_path) -> tuple[str, str, str]:
    client.post("/api/workspace", json={"path": str(tmp_path / "ws")})
    pid = client.post("/api/projects", json={"name": "引用闭环项目"}).json()["id"]
    ws = client.post(f"/api/projects/{pid}/workspace").json()
    return pid, ws["office_document_id"], str(tmp_path / "ws")


def _import_material(client, tmp_path) -> str:
    zp = tmp_path / "素材包.zip"
    with zipfile.ZipFile(zp, "w") as z:
        z.writestr("人员简历.docx", _docx(["张三 简历", "高级工程师 8年"]))
    r = client.post("/api/workspace/imports", json={"paths": [str(zp)]})
    batch_id = r.json()["batch_id"]
    item = client.get(f"/api/workspace/imports/{batch_id}").json()["items"][0]
    client.post(f"/api/workspace/imports/{batch_id}/confirm", json={"decisions": [
        {"item_id": item["id"], "kind": "material", "material_cat": "人员简历"},
    ]})
    return item["id"] if False else client.get("/api/materials").json()["items"][0]["id"]


def test_reference_insert_chain(client, tmp_path):
    """C2 主链：插入 → C-3 新版本 + 台账(active) + 素材自动建档 + 处理器追平登记。"""
    pid, ws_oid, ws_root = _setup_ws_and_project(client, tmp_path)
    # 给工作稿写入带章节的版本 v2（模拟历史工作稿内容）
    from app.models import OfficeDocument
    from app.services.office_document_svc import (
        OfficeDocumentService,
        OfficeDocumentTableAdapter,
    )

    s = _db()
    try:
        od = s.get(OfficeDocument, ws_oid)
        OfficeDocumentService(session=s, adapter=OfficeDocumentTableAdapter(s)).create_version(
            od, _docx(["#1 第三章 技术要求", "项目负责人要求：高级工程师。"]), source="upload")
        s.commit()
    finally:
        s.close()

    material_id = _import_material(client, tmp_path)
    node = client.get(f"/api/office-documents/{ws_oid}/outline").json()["nodes"][0]
    r = client.post("/api/workspace/references/insert", json={
        "oid": ws_oid, "node_id": node["id"], "material_id": material_id,
        "purpose": "项目负责人", "project_id": pid,
    })
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["writeback"] in ("ok", "pending_writeback")

    refs = client.get(f"/api/workspace/references?target_oid={ws_oid}").json()
    assert refs["items"] and refs["items"][0]["status"] in ("active", "drifted")
    assert refs["items"][0]["purpose"] == "项目负责人"
    # C-3 链复用：workspace_mutations 有记录
    s = _db()
    try:
        n = s.execute(__import__("sqlalchemy").text(
            "SELECT count(*) FROM workspace_mutations WHERE workspace_office_document_id=:i"),
            {"i": ws_oid}).scalar()
        assert n >= 1
    finally:
        s.close()


def test_reference_drift_and_relocate(client, tmp_path):
    pid, ws_oid, ws_root = _setup_ws_and_project(client, tmp_path)
    from app.models import OfficeDocument
    from app.services.office_document_svc import (
        OfficeDocumentService,
        OfficeDocumentTableAdapter,
    )

    s = _db()
    try:
        od = s.get(OfficeDocument, ws_oid)
        svc = OfficeDocumentService(session=s, adapter=OfficeDocumentTableAdapter(s))
        svc.create_version(od, _docx(["#1 第三章 技术要求", "负责人要求。"]), source="upload")
        s.commit()
    finally:
        s.close()
    material_id = _import_material(client, tmp_path)
    node = client.get(f"/api/office-documents/{ws_oid}/outline").json()["nodes"][0]
    client.post("/api/workspace/references/insert", json={
        "oid": ws_oid, "node_id": node["id"], "material_id": material_id,
        "purpose": "项目负责人", "project_id": pid,
    })

    # 外部演进：新版本删除该章节 → 引用丢失
    s = _db()
    try:
        od = s.get(OfficeDocument, ws_oid)
        OfficeDocumentService(session=s, adapter=OfficeDocumentTableAdapter(s)).create_version(
            od, _docx(["#1 第九章 其他内容", "无关正文。"]), source="upload")
        s.commit()
    finally:
        s.close()
    refs = client.get(f"/api/workspace/references?target_oid={ws_oid}").json()
    assert refs["items"][0]["status"] == "lost", refs["items"]

    # 重定位到现存章节 → 恢复正常
    nodes = client.get(f"/api/office-documents/{ws_oid}/outline").json()["nodes"]
    r = client.post(f"/api/workspace/references/{refs['items'][0]['id']}/relocate",
                    json={"node_id": nodes[0]["id"]})
    assert r.status_code == 200 and r.json()["ok"] is True
    refs2 = client.get(f"/api/workspace/references?target_oid={ws_oid}").json()
    assert refs2["items"][0]["status"] == "active"


def test_writeback_candidate_on_locked_file(client, tmp_path, monkeypatch):
    """C4：目标投影被占用 → 候选落 中转/待写回，重试成功后清理。"""
    pid, ws_oid, ws_root = _setup_ws_and_project(client, tmp_path)
    from app.models import OfficeDocument
    from app.services.office_document_svc import (
        OfficeDocumentService,
        OfficeDocumentTableAdapter,
    )

    s = _db()
    try:
        od = s.get(OfficeDocument, ws_oid)
        OfficeDocumentService(session=s, adapter=OfficeDocumentTableAdapter(s)).create_version(
            od, _docx(["#1 第三章 技术要求", "负责人要求。"]), source="upload")
        s.commit()
    finally:
        s.close()
    material_id = _import_material(client, tmp_path)
    node = client.get(f"/api/office-documents/{ws_oid}/outline").json()["nodes"]
    node = node[0] if node else None
    assert node is not None

    import os as _os

    real_replace = _os.replace

    def _locked(src, dst, *a, **k):
        if str(dst).endswith(".docx") and "投标文件" in str(dst):
            raise PermissionError(32, "文件被占用")
        return real_replace(src, dst, *a, **k)

    monkeypatch.setattr(_os, "replace", _locked)
    r = client.post("/api/workspace/references/insert", json={
        "oid": ws_oid, "node_id": node["id"], "material_id": material_id,
        "purpose": "项目负责人", "project_id": pid,
    })
    assert r.status_code == 201
    assert r.json()["writeback"] == "pending_writeback"
    assert (tmp_path / "ws" / "中转" / "待写回").is_dir()

    # 解除占用 → 重试写回成功
    monkeypatch.setattr(_os, "replace", real_replace)
    ref_id = r.json()["reference_id"]
    r2 = client.post(f"/api/workspace/references/{ref_id}/writeback")
    assert r2.status_code == 200 and r2.json()["ok"] is True
    assert (tmp_path / "ws" / "项目" / "引用闭环项目" / "投标文件" / "引用闭环项目.docx").is_file()
