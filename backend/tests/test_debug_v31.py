"""临时调试：append_end 输出字节（定位后删除）。"""
from __future__ import annotations

import io
import os
import zipfile

import pytest
from docx import Document as Docx
from fastapi.testclient import TestClient
from sqlalchemy.orm import sessionmaker


@pytest.fixture()
def client(tmp_path, monkeypatch):
    os.environ["ACCESS_PASSWORD"] = ""
    os.environ["DATA_DIR"] = str(tmp_path / "data")
    os.environ["SCAN_INTERVAL_SECONDS"] = "3600"
    os.environ["BIDVAULT_LOCAL_AUTH"] = "0"
    import app.config as cfg
    import app.db as appdb

    monkeypatch.setattr(cfg, "_SETTINGS", None)
    monkeypatch.setattr(appdb, "_Engine", None)
    monkeypatch.setattr(appdb, "_SessionLocal", None)
    from app.db import get_session, make_engine
    from app.main import app

    monkeypatch.setattr(cfg, "_SETTINGS", None)
    engine = make_engine(cfg.get_settings().db_path)
    factory = sessionmaker(bind=engine, expire_on_commit=False)

    def _make():
        s = factory()
        try:
            yield s
        finally:
            s.close()

    app.dependency_overrides[get_session] = _make
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def _docx(lines):
    buf = io.BytesIO()
    d = Docx()
    for ln in lines:
        if ln.startswith("#1 "):
            d.add_heading(ln[3:], level=1)
        else:
            d.add_paragraph(ln)
    d.save(buf)
    return buf.getvalue()


def test_debug_append_end(client, tmp_path):
    c = client
    c.post("/api/workspace", json={"path": str(tmp_path / "ws")})
    pid = c.post("/api/projects", json={"name": "V31 项目"}).json()["id"]
    ws = c.post(f"/api/projects/{pid}/workspace").json()
    ws_oid = ws["office_document_id"]
    from app.config import get_settings
    from app.db import get_engine
    from app.models import OfficeDocument, OfficeDocumentVersion
    from app.services.office_document_svc import (
        OfficeDocumentService,
        OfficeDocumentTableAdapter,
    )

    s = sessionmaker(bind=get_engine(), expire_on_commit=False)()
    od = s.get(OfficeDocument, ws_oid)
    svc = OfficeDocumentService(session=s, adapter=OfficeDocumentTableAdapter(s))
    svc.create_version(od, _docx(["#1 第三章 技术要求", "正文。"]), source="upload")
    s.commit()
    s.close()

    zp = tmp_path / "素材.zip"
    with zipfile.ZipFile(zp, "w") as z:
        z.writestr("人员简历-尾部素材.docx", _docx(["尾部素材 内容"]))
    r = c.post("/api/workspace/imports", json={"paths": [str(zp)]})
    batch = r.json()["batch_id"]
    items = c.get(f"/api/workspace/imports/{batch}").json()["items"]
    c.post(f"/api/workspace/imports/{batch}/confirm", json={"decisions": [
        {"item_id": items[0]["id"], "kind": "material", "material_cat": "人员简历"}]})
    mat = c.get("/api/materials").json()["items"][0]
    node = c.get(f"/api/office-documents/{ws_oid}/outline").json()["nodes"][0]
    r = c.post("/api/workspace/references/insert", json={
        "oid": ws_oid, "node_id": node["id"], "material_id": mat["id"],
        "purpose": "项目负责人", "insert_mode": "append_end"})
    print("INSERT:", r.status_code, r.text[:200])
    s = sessionmaker(bind=get_engine(), expire_on_commit=False)()
    od = s.get(OfficeDocument, ws_oid)
    v = s.get(OfficeDocumentVersion, od.current_version_id)
    from app.config import get_settings
    from app.db import get_engine

    data = (get_settings().files_dir / v.file_path).read_bytes()
    print("NEW VERSION head:", data[:8], "size:", len(data))
    s.close()
    assert True
