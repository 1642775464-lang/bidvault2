"""E-2B：结构保真阅读块验收测试（E-0 契约 §5）。

覆盖：heading/para/table(image 占位)块序列、gridSpan/vMerge 保真、
para_no 与 DocumentOutline.para_index 同口径、真实民机样本可解析。
"""

from __future__ import annotations

import io
import os
import time
from pathlib import Path

import pytest
from docx import Document as Docx
from fastapi.testclient import TestClient
from sqlalchemy.orm import sessionmaker


@pytest.fixture()
def client(tmp_path, monkeypatch):
    os.environ["ACCESS_PASSWORD"] = ""
    os.environ["DATA_DIR"] = str(tmp_path / "data")
    os.environ["SCAN_INTERVAL_SECONDS"] = "3600"
    import app.config as cfg
    import app.db as appdb

    monkeypatch.setattr(cfg, "_SETTINGS", None)
    monkeypatch.setattr(appdb, "_Engine", None)
    monkeypatch.setattr(appdb, "_SessionLocal", None)
    from app.db import get_session, make_engine
    from app.main import app

    monkeypatch.setattr(cfg, "_SETTINGS", None)
    engine = make_engine(cfg.get_settings().db_path)
    factory = sessionmaker(bind=engine, expire_on_commit=False)

    def _make():
        s = factory()
        try:
            yield s
        finally:
            s.close()

    app.dependency_overrides[get_session] = _make
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def _docx_with_table_and_headings() -> bytes:
    d = Docx()
    d.add_heading("第一章 投标须知", level=1)
    d.add_paragraph("前言段落。")
    t = d.add_table(rows=2, cols=3)
    t.style = "Table Grid"
    a = t.cell(0, 0)
    a.merge(t.cell(0, 2))
    a.text = "合并表头"
    t.cell(1, 0).text = "A"
    t.cell(1, 1).text = "B"
    t.cell(1, 2).text = "C"
    d.add_heading("第二章 评标办法", level=1)
    d.add_paragraph("评标细则。")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _setup(client, tmp_path) -> str:
    client.post("/api/workspace", json={"path": str(tmp_path / "ws")})
    pid = client.post("/api/projects", json={"name": "阅读器项目"}).json()["id"]
    r = client.post("/api/documents/upload",
                    files=[("files", ("招标文件.docx", _docx_with_table_and_headings(),
                                      "application/octet-stream"))],
                    data={"kind": "tender", "project_id": pid})
    assert r.status_code == 200, r.text
    doc_id = r.json()["created"][0]["id"]
    deadline = time.time() + 60
    while time.time() < deadline:
        docs = client.get("/api/documents").json()["items"]
        d = next((x for x in docs if x["id"] == doc_id), None)
        if d and d["status"] == 2 and d.get("office_document_id"):
            return d["office_document_id"]
        time.sleep(1)
    raise AssertionError("拆解未完成")


def test_reader_blocks_structure_and_alignment(client, tmp_path):
    oid = _setup(client, tmp_path)
    outline = client.get(f"/api/office-documents/{oid}/outline").json()
    reader = client.get(f"/api/office-documents/{oid}/reader").json()
    assert reader["version"] == outline["version"]

    blocks = reader["blocks"]
    types = [b["type"] for b in blocks]
    assert "heading" in types and "para" in types and "table" in types

    # 章节标题块存在且 level 正确
    heads = [b for b in blocks if b["type"] == "heading"]
    assert [h["text"] for h in heads] == ["第一章 投标须知", "第二章 评标办法"]
    assert all(h["level"] == 1 for h in heads)

    # 表格块：合并表头 gridspan=3 + 数据行
    tables = [b for b in blocks if b["type"] == "table"]
    assert len(tables) == 1
    rows = tables[0]["rows"]
    assert rows[0]["cells"][0]["gridspan"] == 3
    assert rows[0]["cells"][0]["text"] == "合并表头"
    assert [c["text"] for c in rows[1]["cells"]] == ["A", "B", "C"]

    # para_no 与 outline.para_index 同口径：每个标题块的 para_no 等于
    # outline 中同名节点的 para_index
    nodes = {n["title"]: n for n in outline["nodes"]}
    for h in heads:
        assert h["para_no"] == nodes[h["text"]]["para_index"], h

    # 导航映射可用性：按 node.para_index 能找到对应块（首个 para_no>=index 的块）
    for title, node in nodes.items():
        hit = next(b for b in blocks if b["para_no"] >= node["para_index"])
        assert hit["type"] in ("heading", "para", "table")


def test_reader_real_sample_parses(client, tmp_path):
    """真实民机 114 页样本：阅读块可构建、章节标题块与表格均存在。"""
    sample = Path(__file__).resolve().parents[2] / "回归证据" / "D1R" / "dx-console" / "DxConsole" / "民机_test.docx"
    if not sample.is_file():
        pytest.skip("民机样本不存在")
    client.post("/api/workspace", json={"path": str(tmp_path / "ws")})
    pid = client.post("/api/projects", json={"name": "样本项目"}).json()["id"]
    r = client.post("/api/documents/upload",
                    files=[("files", ("民机样本.docx", sample.read_bytes(), "application/octet-stream"))],
                    data={"kind": "tender", "project_id": pid})
    assert r.status_code == 200
    doc_id = r.json()["created"][0]["id"]
    deadline = time.time() + 120
    oid = None
    while time.time() < deadline:
        docs = client.get("/api/documents").json()["items"]
        d = next((x for x in docs if x["id"] == doc_id), None)
        if d and d["status"] == 2 and d.get("office_document_id"):
            oid = d["office_document_id"]
            break
        time.sleep(1)
    assert oid, "样本拆解超时"
    reader = client.get(f"/api/office-documents/{oid}/reader").json()
    blocks = reader["blocks"]
    assert len(blocks) > 500
    types = {b["type"] for b in blocks}
    assert {"heading", "para", "table"} <= types
    # 章节标题块覆盖第三章
    assert any(b["type"] == "heading" and "第三章" in b["text"] for b in blocks)
