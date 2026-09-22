"""阶段六 C-3b：OOXML 包级插入引擎测试（16 项矩阵）+ 服务级真引擎集成。

- 引擎级（快速、无 HTTP）：段落/表格/gridSpan/vMerge/图片/超链接/样式冲突/numbering 冲突/
  media 冲突/sectPr 保持/结构拒绝/XXE 拒绝
- 服务级（真引擎）：stale base 409、活跃编辑 423、失败零副作用、Source/Material 文件不变
"""

from __future__ import annotations

import base64
import io
import os
import re
import time
import zipfile

import pytest
from docx import Document as Docx
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches
from fastapi.testclient import TestClient
from lxml import etree
from sqlalchemy.orm import sessionmaker

from app.db import get_session, make_engine
from app.main import app
from app.models import (
    AuditLog,
    DocumentVersion,
    OfficeDocument,
    OfficeDocumentVersion,
    WorkspaceDocument,
    WorkspaceMutation,
)
from app.services.ooxml_insert import OOXMLInsertEngine
from app.services.xml_safety import UnsafeXmlError, safe_parse_xml

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@pytest.fixture()
def client(tmp_path, monkeypatch):
    os.environ["ACCESS_PASSWORD"] = ""
    os.environ["SOURCE_MAINTENANCE_MODE"] = "0"
    os.environ["DATA_DIR"] = str(tmp_path / "data")
    os.environ["SCAN_INTERVAL_SECONDS"] = "3600"
    import app.config as cfg
    import app.db as appdb

    monkeypatch.setattr(cfg, "_SETTINGS", None)
    monkeypatch.setattr(appdb, "_Engine", None)
    monkeypatch.setattr(appdb, "_SessionLocal", None)
    engine = make_engine(cfg.get_settings().db_path)
    app.dependency_overrides[get_session] = _session_factory(engine)
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def _session_factory(engine):
    factory = sessionmaker(bind=engine, expire_on_commit=False)

    def _make():
        s = factory()
        try:
            yield s
        finally:
            s.close()

    return _make


def _db_session():
    import app.config as cfg

    return sessionmaker(bind=make_engine(cfg.get_settings().db_path), expire_on_commit=False)()


def _blank_ws() -> bytes:
    buf = io.BytesIO()
    d = Docx()
    d.add_heading("测试 投标工作稿", level=0)
    d.add_paragraph("正文一段")
    d.save(buf)
    return buf.getvalue()


def _tiny_png() -> bytes:
    return base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQ"
        "AAAABJRU5ErkJggg==")


def _body_of(docx_bytes: bytes):
    z = zipfile.ZipFile(io.BytesIO(docx_bytes))
    return safe_parse_xml(z.read("word/document.xml")).find(_W + "body")


def _texts(body) -> list[str]:
    return ["".join(t.text or "" for t in p.iter(_W + "t")) for p in body.findall(_W + "p")]


def _engine() -> OOXMLInsertEngine:
    return OOXMLInsertEngine()


def _mutate(ws: bytes, mat: bytes) -> bytes:
    return _engine().mutate(None, ws, mat)  # type: ignore[arg-type]


# ---------------- 1-2：段落 / 表格 ----------------


def test_plain_paragraph_append():
    mat_buf = io.BytesIO()
    m = Docx()
    m.add_paragraph("素材一")
    m.add_paragraph("素材二")
    m.save(mat_buf)
    out = _mutate(_blank_ws(), mat_buf.getvalue())
    body = _body_of(out)
    assert _texts(body) == ["测试 投标工作稿", "正文一段", "素材一", "素材二"]
    assert body[-1].tag == _W + "sectPr"  # sectPr 仍是末尾


def test_plain_table_append():
    mat_buf = io.BytesIO()
    m = Docx()
    t = m.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    t.cell(0, 0).text = "A1"
    t.cell(1, 1).text = "B2"
    m.save(mat_buf)
    out = _mutate(_blank_ws(), mat_buf.getvalue())
    d = Docx(io.BytesIO(out))
    assert len(d.tables) == 1
    assert d.tables[0].cell(0, 0).text == "A1" and d.tables[0].cell(1, 1).text == "B2"


# ---------------- 3-4：gridSpan / vMerge ----------------


def test_gridspan_preserved():
    mat_buf = io.BytesIO()
    m = Docx()
    t = m.add_table(rows=2, cols=3)
    t.cell(0, 0).merge(t.cell(0, 2)).text = "跨三列"
    m.save(mat_buf)
    out = _mutate(_blank_ws(), mat_buf.getvalue())
    body = _body_of(out)
    assert sum(1 for _ in body.iter(_W + "gridSpan")) == 1
    d = Docx(io.BytesIO(out))
    merged = d.tables[-1].cell(0, 0)
    assert merged.text == "跨三列" and d.tables[-1].cell(0, 1)._tc is merged._tc


def test_vmerge_preserved():
    mat_buf = io.BytesIO()
    m = Docx()
    t = m.add_table(rows=3, cols=2)
    t.cell(0, 0).merge(t.cell(1, 0)).text = "跨两行"
    m.save(mat_buf)
    out = _mutate(_blank_ws(), mat_buf.getvalue())
    d = Docx(io.BytesIO(out))
    merged = d.tables[-1].cell(0, 0)
    assert merged.text == "跨两行" and d.tables[-1].cell(1, 0)._tc is merged._tc


# ---------------- 5-6：图片 / 超链接 relationship 重映射 ----------------


def test_image_relationship_remap():
    mat_buf = io.BytesIO()
    m = Docx()
    m.add_picture(io.BytesIO(_tiny_png()), width=Inches(1))
    m.save(mat_buf)
    orig_rids = set(re.findall(r'Id="(rId\d+)"',
                               zipfile.ZipFile(io.BytesIO(mat_buf.getvalue())).read(
                                   "word/_rels/document.xml.rels").decode()))
    out = _mutate(_blank_ws(), mat_buf.getvalue())
    z = zipfile.ZipFile(io.BytesIO(out))
    media = [n for n in z.namelist() if n.startswith("word/media/")]
    assert media, "media part 未落盘"
    rels = z.read("word/_rels/document.xml.rels").decode()
    doc = z.read("word/document.xml").decode()
    for rid in orig_rids:
        assert f'r:embed="{rid}"' not in doc and f'r:id="{rid}"' not in doc, "复用了 Material 原 rId"
    for name in media:
        part_rel = re.search(rf'Target="media/{name.split("/")[-1]}"', rels)
        assert part_rel, f"media {name} 缺 relationship"
    # r:embed 指向的 rId 在 rels 中存在
    for rid in set(re.findall(r'r:embed="(rId\d+)"', doc)):
        assert f'Id="{rid}"' in rels


def test_hyperlink_remap():
    mat_buf = io.BytesIO()
    m = Docx()
    p = m.add_paragraph("带链接")
    rid = m.part.relate_to(
        "https://example.com/page",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    p._p.append(parse_xml(
        f'<w:hyperlink {nsdecls("w", "r")} r:id="{rid}">'
        f'<w:r><w:t>链接</w:t></w:r></w:hyperlink>'))
    m.save(mat_buf)
    out = _mutate(_blank_ws(), mat_buf.getvalue())
    z = zipfile.ZipFile(io.BytesIO(out))
    rels = z.read("word/_rels/document.xml.rels").decode()
    doc = z.read("word/document.xml").decode()
    assert f'r:id="{rid}"' not in doc, "复用了 Material 原 rId"
    m2 = re.search(r'<Relationship[^>]*Target="https://example.com/page"[^>]*/>', rels)
    assert m2 and 'TargetMode="External"' in m2.group(0)
    new_rid = re.search(r'Id="(rId\d+)"[^>]*Target="https://example.com/page"', rels).group(1)
    assert f'w:hyperlink r:id="{new_rid}"' in doc or f'r:id="{new_rid}"' in doc


# ---------------- 7：style ID 冲突（同 id 不同定义 → 改名复制） ----------------


def test_style_id_conflict_renamed():
    ws_buf = io.BytesIO()
    w = Docx()
    st = w.styles.add_style("Custom1", WD_STYLE_TYPE.PARAGRAPH)
    st.font.size = None
    w.add_paragraph("用样式")
    w.save(ws_buf)
    mat_buf = io.BytesIO()
    m = Docx()
    st2 = m.styles.add_style("Custom1", WD_STYLE_TYPE.PARAGRAPH)
    st2.font.size = __import__("docx").shared.Pt(28)
    p = m.add_paragraph("素材样式段")
    p.style = m.styles["Custom1"]
    m.save(mat_buf)

    out = _mutate(ws_buf.getvalue(), mat_buf.getvalue())
    z = zipfile.ZipFile(io.BytesIO(out))
    styles = z.read("word/styles.xml").decode()
    doc = z.read("word/document.xml").decode()
    renamed = re.findall(r'w:styleId="(Material_Custom1_[0-9a-f]{8})"', styles)
    assert renamed, "同 id 不同定义未改名复制"
    assert renamed[0] in doc, "插入内容未重写为新 styleId"
    # Workspace 原有 Custom1 定义保持
    base_styles = zipfile.ZipFile(io.BytesIO(ws_buf.getvalue())).read("word/styles.xml").decode()
    assert base_styles.count('w:styleId="Custom1"') == 1
    assert styles.count('w:styleId="Custom1"') == 1, "Workspace 原 Custom1 被污染"


# ---------------- 8：numbering ID 冲突（一律新 id） ----------------


def _numbering_xml() -> bytes:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:numbering xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:abstractNum w:abstractNumId="0"><w:lvl w:ilvl="0">'
        '<w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="•"/>'
        '</w:lvl></w:abstractNum>'
        '<w:num w:numId="3"><w:abstractNumId w:val="0"/></w:num>'
        '</w:numbering>'
    ).encode()


def _with_numbering(docx_bytes: bytes) -> bytes:
    """把模板 numbering.xml 替换为已知定义（numId=3 → bullet abstractNum 0）。"""
    src = zipfile.ZipFile(io.BytesIO(docx_bytes))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in src.infolist():
            zo.writestr(item, _numbering_xml() if item.filename == "word/numbering.xml"
                        else src.read(item.filename))
    return out.getvalue()


def test_numbering_id_conflict_reassigned():
    mat_buf = io.BytesIO()
    m = Docx()
    p = m.add_paragraph("列表项")
    num_pr = p._p.get_or_add_pPr()
    num_pr.append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="3"/></w:numPr>'))
    m.save(mat_buf)
    mat_bytes = _with_numbering(mat_buf.getvalue())
    ws_bytes = _blank_ws()

    ws_num_ids = set(re.findall(r'<w:num w:numId="(\d+)"',
                                zipfile.ZipFile(io.BytesIO(ws_bytes)).read("word/numbering.xml").decode()))
    out = _mutate(ws_bytes, mat_bytes)
    z = zipfile.ZipFile(io.BytesIO(out))
    numxml = z.read("word/numbering.xml").decode()
    doc = z.read("word/document.xml").decode()

    used = set(re.findall(r'<w:numId w:val="(\d+)"', doc))
    assert used, "插入段落未引用 numbering"
    assert not (used & ws_num_ids), f"复用了 Workspace 原 numId：{used & ws_num_ids}"
    # 新 num 的 abstract 携带素材的 bullet 定义
    new_ids = used - {nid for nid in used if nid in ws_num_ids}
    for nid in new_ids:
        abs_ref = re.search(rf'<w:num w:numId="{nid}"><w:abstractNumId w:val="(\d+)"/>', numxml)
        assert abs_ref, f"新 num {nid} 缺 abstractNumId"
        assert 'w:numFmt w:val="bullet"' in numxml, "bullet 定义未复制"


# ---------------- 9：media 文件名冲突（不覆盖） ----------------


def test_media_filename_conflict_no_overwrite():
    ws_buf = io.BytesIO()
    w = Docx()
    w.add_picture(io.BytesIO(_tiny_png()), width=Inches(1))
    w.add_paragraph("工作稿已有图")
    w.save(ws_buf)
    mat_buf = io.BytesIO()
    m = Docx()
    m.add_picture(io.BytesIO(_tiny_png() + b"\x00\x01"), width=Inches(1))  # 不同内容
    m.save(mat_buf)

    out = _mutate(ws_buf.getvalue(), mat_buf.getvalue())
    z = zipfile.ZipFile(io.BytesIO(out))
    media = sorted(n for n in z.namelist() if n.startswith("word/media/"))
    assert len(media) == 2, f"media 数量异常：{media}"
    blobs = {z.read(n) for n in media}
    assert len(blobs) == 2, "media 内容被覆盖"


# ---------------- 10：Workspace final sectPr 保持 ----------------


def test_final_sectpr_preserved():
    ws_bytes = _blank_ws()
    mat_buf = io.BytesIO()
    m = Docx()
    m.add_paragraph("插入内容")
    m.save(mat_buf)

    ws_sectpr = etree.tostring(_body_of(ws_bytes)[-1])
    out = _mutate(ws_bytes, mat_buf.getvalue())
    body = _body_of(out)
    assert body[-1].tag == _W + "sectPr"
    assert etree.tostring(body[-1]) == ws_sectpr, "final sectPr 被改动"
    # 插入位置在 sectPr 之前：原有段落仍在素材段落之前
    texts = _texts(body)
    assert texts[:2] == ["测试 投标工作稿", "正文一段"] and texts[-1] == "插入内容"


# ---------------- 11：Material 中间 section → 422 ----------------


def test_material_mid_section_rejected():
    mat_buf = io.BytesIO()
    m = Docx()
    m.add_paragraph("分节前的段落")
    m.save(mat_buf)
    d = Docx(io.BytesIO(mat_buf.getvalue()))
    p = d.paragraphs[0]
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:sectPr {nsdecls("w")}><w:pgSz w:w="11906" w:h="16838"/></w:sectPr>'))
    buf2 = io.BytesIO()
    d.save(buf2)

    with pytest.raises(Exception) as ei:
        _mutate(_blank_ws(), buf2.getvalue())
    assert ei.value.code == "UNSUPPORTED_MATERIAL_STRUCTURE"
    assert "section_break" in ei.value.features


# ---------------- 16：XXE / DOCTYPE 拒绝 ----------------


def test_xxe_doctype_rejected():
    evil_doc = (
        b'<?xml version="1.0"?><!DOCTYPE foo [<!ENTITY xxe SYSTEM "file:///C:/Windows/win.ini">]>'
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        b'<w:body><w:p><w:r><w:t>&xxe;</w:t></w:r></w:p></w:body></w:document>'
    )
    evil_zip = io.BytesIO()
    with zipfile.ZipFile(evil_zip, "w") as zo:
        for name in ("[Content_Types].xml", "word/document.xml", "word/_rels/document.xml.rels"):
            zo.writestr(name, b"<x/>" if name != "word/document.xml" else evil_doc)
    with pytest.raises(UnsafeXmlError):
        _mutate(_blank_ws(), evil_zip.getvalue())
    from app.services.xml_safety import safe_parse_xml as sp

    with pytest.raises(UnsafeXmlError):
        sp(evil_doc)
    # 正常 XML 不受影响
    sp(b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
       b'<w:body><w:p/></w:body></w:document>')


# ---------------- 服务级集成（真引擎）：12-15 ----------------


def _setup_word_material(client) -> tuple[str, str]:
    buf = io.BytesIO()
    d = Docx()
    d.add_paragraph("投标人须知前附表")
    d.add_paragraph("递交截止时间：2026-10-01 09:00")
    d.save(buf)
    pid = client.post("/api/projects", json={"name": "C3b 素材项目"}).json()["id"]
    r = client.post(
        "/api/documents/upload",
        files=[("files", ("源文件.docx", buf.getvalue(), "application/octet-stream"))],
        data={"kind": "tender", "project_id": pid, "is_secret": "false"},
    )
    doc_id = r.json()["created"][0]["id"]
    deadline = time.time() + 30
    while time.time() < deadline:
        docs = client.get("/api/documents").json()["items"]
        if next(x for x in docs if x["id"] == doc_id)["status"] == 2:
            break
        time.sleep(0.2)
    client.get(f"/api/documents/{doc_id}/editor-config")
    m = client.post("/api/materials", json={"title": "C3b 素材", "cat": "投标格式", "type": "文本"}).json()
    r2 = client.post(
        f"/api/materials/{m['id']}/office-document",
        json={"source_document_id": doc_id, "child_range": [0, 1], "comment": "提取"},
    )
    assert r2.status_code == 200, r2.text
    return m["id"], r2.json()["office_document_id"]


def _setup_workspace(client) -> str:
    pid = client.post("/api/projects", json={"name": "C3b 工作稿项目"}).json()["id"]
    r = client.post(f"/api/projects/{pid}/workspace")
    return r.json()["office_document_id"]


def _service(_client, engine=None, probe=None):
    from app.models import OfficeDocument, OfficeDocumentVersion
    from app.services.office_document_svc import dispatch_post_save
    from app.services.workspace_mutation import WorkspaceMutationService

    def _sync_dispatch(oid: str, version_id: str) -> None:
        s = _db_session()
        try:
            od = s.get(OfficeDocument, oid)
            v = s.get(OfficeDocumentVersion, version_id)
            if od is not None and v is not None:
                dispatch_post_save(od, v)
        finally:
            s.close()

    return WorkspaceMutationService(
        session=_db_session(),
        active_edit_probe=probe or (lambda _oid: False),
        engine=engine if engine is not None else OOXMLInsertEngine(),
        dispatch_fn=_sync_dispatch,
    )


def _insert_req(ws_oid: str, base_vid: str, mid: str, mod_oid: str, mod_vid: str, **kw):
    from app.services.workspace_mutation import MaterialInsertRequest

    return MaterialInsertRequest(
        workspace_office_document_id=ws_oid, base_workspace_version_id=base_vid,
        material_id=mid, material_office_document_id=mod_oid,
        material_version_id=mod_vid, **kw)


def _version_id(oid: str, version: int) -> str:
    s = _db_session()
    try:
        v = (s.query(OfficeDocumentVersion)
             .filter(OfficeDocumentVersion.office_document_id == oid,
                     OfficeDocumentVersion.version == version).one())
        return v.id
    finally:
        s.close()


def test_service_real_engine_stale_base_and_active_edit(client):
    """12/13：真引擎下 stale base 仍 409、活跃编辑仍 423。"""
    from app.services.workspace_mutation import WorkspaceMutationError

    mid, mod_oid = _setup_word_material(client)
    ws_oid = _setup_workspace(client)
    v1 = _version_id(ws_oid, 1)
    mat_vid = _version_id(mod_oid, 1)
    svc = _service(client)

    r = svc.insert_material(_insert_req(ws_oid, v1, mid, mod_oid, mat_vid))
    assert r["version"] == 2
    with pytest.raises(WorkspaceMutationError) as ei:
        svc.insert_material(_insert_req(ws_oid, v1, mid, mod_oid, mat_vid))
    assert ei.value.code == "WORKSPACE_VERSION_CONFLICT"
    svc2 = _service(client, probe=lambda _oid: True)
    with pytest.raises(WorkspaceMutationError) as ei2:
        svc2.insert_material(_insert_req(ws_oid, _version_id(ws_oid, 2), mid, mod_oid, mat_vid))
    assert ei2.value.code == "WORKSPACE_EDITING"


def test_service_unsupported_structure_zero_side_effects(client):
    """14：结构不支持 → 422 且零副作用（无新版本/无审计/current 不变）。"""
    from app.services.workspace_mutation import UnsupportedMaterialStructureError

    mid, mod_oid = _setup_word_material(client)
    ws_oid = _setup_workspace(client)
    v1 = _version_id(ws_oid, 1)
    mat_vid = _version_id(mod_oid, 1)

    # 给素材 v1 注入脚注引用 → 不支持结构
    import app.config as cfg

    files_dir = cfg.get_settings().files_dir
    s = _db_session()
    try:
        v = s.get(OfficeDocumentVersion, mat_vid)
        data = (files_dir / v.file_path).read_bytes()
        z = zipfile.ZipFile(io.BytesIO(data))
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
            for item in z.infolist():
                if item.filename == "word/document.xml":
                    doc = z.read(item.filename).decode()
                    doc = doc.replace("<w:p>", '<w:p><w:r><w:footnoteReference w:id="2"/></w:r>', 1)
                    zo.writestr(item, doc)
                else:
                    zo.writestr(item, z.read(item.filename))
        v.file_path = f"material_evil_{int(time.time())}.docx"
        (files_dir / v.file_path).write_bytes(out.getvalue())
        s.commit()
    finally:
        s.close()

    before_ver = _version_id
    svc = _service(client)
    with pytest.raises(UnsupportedMaterialStructureError) as ei:
        svc.insert_material(_insert_req(ws_oid, v1, mid, mod_oid, mat_vid))
    assert "footnote_endnote" in ei.value.features

    s = _db_session()
    try:
        ver_count = (s.query(OfficeDocumentVersion)
                     .filter(OfficeDocumentVersion.office_document_id == ws_oid).count())
        audit_count = s.query(WorkspaceMutation).count()
        current = s.get(OfficeDocument, ws_oid).current_version_id
    finally:
        s.close()
    assert ver_count == 1 and audit_count == 0 and current == v1
    del before_ver


def test_service_source_material_files_unchanged(client):
    """15：完整插入后 Source 拆解文件与 Material 版本文件 hash 完全不变。"""
    mid, mod_oid = _setup_word_material(client)
    ws_oid = _setup_workspace(client)
    v1 = _version_id(ws_oid, 1)
    mat_vid = _version_id(mod_oid, 1)

    s = _db_session()
    try:
        src_files = [(v.id, v.file_hash) for v in s.query(DocumentVersion).all()]
        mat_files = [(v.id, v.file_hash) for v in
                     s.query(OfficeDocumentVersion)
                     .filter(OfficeDocumentVersion.office_document_id == mod_oid).all()]
        audit_before = s.query(AuditLog).filter(AuditLog.action == "oo_office_version_saved").count()
    finally:
        s.close()

    r = _service(client).insert_material(_insert_req(ws_oid, v1, mid, mod_oid, mat_vid))
    assert r["version"] == 2

    s = _db_session()
    try:
        assert [(v.id, v.file_hash) for v in s.query(DocumentVersion).all()] == src_files
        assert [(v.id, v.file_hash) for v in
                s.query(OfficeDocumentVersion)
                .filter(OfficeDocumentVersion.office_document_id == mod_oid).all()] == mat_files
        # 未伪造 ONLYOFFICE callback
        assert s.query(AuditLog).filter(
            AuditLog.action == "oo_office_version_saved").count() == audit_before
        audit = s.query(WorkspaceMutation).one()
        assert (audit.workspace_base_version_id, audit.workspace_new_version_id) == (v1, r["version_id"])
        derived = (s.query(WorkspaceDocument)
                   .filter(WorkspaceDocument.office_document_id == ws_oid).one())
    finally:
        s.close()
    assert derived.derived_from_version_id == r["version_id"]
