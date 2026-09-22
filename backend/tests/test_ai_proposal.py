"""阶段六 C-4a：AI 提案契约测试（确定性替身；无真实 AI 调用、无真实 OOXML 写入）。

契约见 阶段六C4a-AI审批契约设计.md：
- AI 内容在 approve 前只落 ai_proposals（禁进 derived_text/content_text/版本表）
- apply = create_version(source="ai_insert")；CAS base==current（失效→expired 409）
- 状态机：rejected/expired/applied 不可应用；重复 apply 不产生第二个版本
- apply 失败 → failed + 零副作用；不写 Source/Material；AI 失败不影响 Workspace
"""

from __future__ import annotations

import io
import json
import os
import time

import pytest
from docx import Document as Docx
from fastapi.testclient import TestClient
from sqlalchemy.orm import sessionmaker

from app.db import get_session, make_engine
from app.main import app
from app.models import (
    AiProposal,
    AuditLog,
    Document,
    DocumentVersion,
    Material,
    OfficeDocument,
    OfficeDocumentVersion,
    WorkspaceDocument,
)


@pytest.fixture()
def client(tmp_path, monkeypatch):
    os.environ["ACCESS_PASSWORD"] = ""
    os.environ["SOURCE_MAINTENANCE_MODE"] = "0"
    os.environ["DATA_DIR"] = str(tmp_path / "data")
    os.environ["SCAN_INTERVAL_SECONDS"] = "3600"
    import app.config as cfg
    import app.db as appdb

    monkeypatch.setattr(cfg, "_SETTINGS", None)
    monkeypatch.setattr(appdb, "_Engine", None)
    monkeypatch.setattr(appdb, "_SessionLocal", None)
    engine = make_engine(cfg.get_settings().db_path)
    app.dependency_overrides[get_session] = _session_factory(engine)
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def _session_factory(engine):
    factory = sessionmaker(bind=engine, expire_on_commit=False)

    def _make():
        s = factory()
        try:
            yield s
        finally:
            s.close()

    return _make


def _db_session():
    import app.config as cfg

    return sessionmaker(bind=make_engine(cfg.get_settings().db_path), expire_on_commit=False)()


def _docx_bytes(content: str) -> bytes:
    buf = io.BytesIO()
    d = Docx()
    for line in content.split("\n"):
        d.add_paragraph(line)
    d.save(buf)
    return buf.getvalue()


def _setup_workspace(client) -> tuple[str, str]:
    """上传 source（供隔离断言用）+ 建项目 + 工作稿。返回 (ws_oid, source_doc_id)。"""
    src = _docx_bytes("招标文件正文：递交截止时间 2026-10-01。")
    pid = client.post("/api/projects", json={"name": "C4a 项目"}).json()["id"]
    r = client.post(
        "/api/documents/upload",
        files=[("files", ("招标.docx", src, "application/octet-stream"))],
        data={"kind": "tender", "project_id": pid, "is_secret": "false"},
    )
    doc_id = r.json()["created"][0]["id"]
    deadline = time.time() + 30
    while time.time() < deadline:
        docs = client.get("/api/documents").json()["items"]
        if next(x for x in docs if x["id"] == doc_id)["status"] == 2:
            break
        time.sleep(0.2)
    ws = client.post(f"/api/projects/{pid}/workspace").json()
    return ws["office_document_id"], doc_id


def _current_vid(ws_oid: str) -> str:
    s = _db_session()
    try:
        return s.get(OfficeDocument, ws_oid).current_version_id
    finally:
        s.close()


def _ver_count(ws_oid: str) -> int:
    s = _db_session()
    try:
        return (s.query(OfficeDocumentVersion)
                .filter(OfficeDocumentVersion.office_document_id == ws_oid).count())
    finally:
        s.close()


class _AiAppendStub:
    """替身 mutator：把提案 text 逐段追加（合法 DOCX 变更；真实 mutator 属 C-4b）。"""

    def __init__(self, fail: bool = False):
        self.fail = fail
        self.calls = 0

    def mutate_append(self, proposal, workspace_bytes: bytes) -> bytes:
        self.calls += 1
        if self.fail:
            raise RuntimeError("AI 写入失败（替身注入）")
        d = Docx(io.BytesIO(workspace_bytes))
        text = json.loads(proposal.content_json)["text"]
        for line in text.split("\n"):
            d.add_paragraph(line)
        out = io.BytesIO()
        d.save(out)
        return out.getvalue()


def _service(_client, mutator=None, probe=None):
    from app.models import OfficeDocument, OfficeDocumentVersion
    from app.services.ai_proposal import AiProposalService
    from app.services.office_document_svc import dispatch_post_save

    def _sync_dispatch(oid: str, version_id: str) -> None:
        s = _db_session()
        try:
            od = s.get(OfficeDocument, oid)
            v = s.get(OfficeDocumentVersion, version_id)
            if od is not None and v is not None:
                dispatch_post_save(od, v)
        finally:
            s.close()

    return AiProposalService(
        session=_db_session(),
        active_edit_probe=probe or (lambda _oid: False),
        mutator=mutator if mutator is not None else _AiAppendStub(),
        dispatch_fn=_sync_dispatch,
    )


def _req(ws_oid: str, base_vid: str, **kw):
    from app.services.ai_proposal import AiProposalRequest

    body = {
        "workspace_office_document_id": ws_oid,
        "base_workspace_version_id": base_vid,
        "title": "AI 建议补充服务承诺",
        "text": "我方承诺提供 7×24 小时响应服务。",
        "model": "glm-4-flash",
        "context": [{"doc": "招标文件", "chapter": "服务要求"}],
    }
    body.update(kw)
    return AiProposalRequest(**body)


def _fingerprint(s) -> tuple:
    docs = tuple((d.id, d.content_text) for d in s.query(Document).all())
    docver = tuple((v.id, v.file_hash) for v in s.query(DocumentVersion).all())
    mats = tuple((m.id, m.title, m.text) for m in s.query(Material).all())
    return (docs, docver, mats)


# ---------------- 1：create draft（round-trip + 未确认内容不落正文） ----------------


def test_create_draft_roundtrip_and_no_body_leak(client):
    ws_oid, _doc_id = _setup_workspace(client)
    base_vid = _current_vid(ws_oid)

    s = _db_session()
    try:
        ver_before = _ver_count(ws_oid)
        derived_before = (s.query(WorkspaceDocument)
                          .filter(WorkspaceDocument.office_document_id == ws_oid)
                          .one().derived_text) or ""
    finally:
        s.close()

    r = _service(client).create_proposal(_req(ws_oid, base_vid))
    assert r["status"] == "draft"

    s = _db_session()
    try:
        p = s.get(AiProposal, r["id"])
        assert p.base_workspace_version_id == base_vid
        assert p.operation == "append_end"
        assert "7×24" in json.loads(p.content_json)["text"]
        assert p.model == "glm-4-flash"
        assert "招标文件" in (p.context_json or "")
        assert p.applied_version_id is None
        # 未确认的 AI 内容不落正文三落点
        assert _ver_count(ws_oid) == ver_before, "draft 阶段产生了版本"
        d_after = (s.query(WorkspaceDocument)
                   .filter(WorkspaceDocument.office_document_id == ws_oid).one())
        assert (d_after.derived_text or "") == derived_before, "derived_text 被改动"
        assert "7×24" not in (d_after.derived_text or "")
        created = s.query(AuditLog).filter(AuditLog.action == "ai_proposal_created").all()
    finally:
        s.close()
    assert created and r["id"][:8] in created[0].detail


# ---------------- 2：approve + apply 全链 ----------------


def test_approve_apply_full_chain(client):
    ws_oid, _doc_id = _setup_workspace(client)
    base_vid = _current_vid(ws_oid)
    svc = _service(client)

    p = svc.create_proposal(_req(ws_oid, base_vid))
    approved = svc.approve_proposal(p["id"])
    assert approved["status"] == "approved"

    applied = svc.apply_proposal(p["id"])
    assert applied["status"] == "applied" and applied["applied_version_id"]

    s = _db_session()
    try:
        v2 = (s.query(OfficeDocumentVersion)
              .filter(OfficeDocumentVersion.office_document_id == ws_oid,
                      OfficeDocumentVersion.version == 2).one())
        current = s.get(OfficeDocument, ws_oid).current_version_id
        row = s.get(AiProposal, p["id"])
        derived = (s.query(WorkspaceDocument)
                   .filter(WorkspaceDocument.office_document_id == ws_oid).one())
        applied_log = s.query(AuditLog).filter(AuditLog.action == "ai_proposal_applied").all()
    finally:
        s.close()
    assert v2.source == "ai_insert"
    assert current == applied["applied_version_id"] == v2.id
    assert row.status == "applied" and row.applied_at is not None
    assert derived.derived_from_version_id == v2.id  # 同步 dispatch：派生追平
    assert applied_log and v2.id in applied_log[0].detail


# ---------------- 3：applied 重复 apply 幂等返回既有结果 ----------------


def test_apply_idempotent_after_applied(client):
    ws_oid, _doc_id = _setup_workspace(client)
    base_vid = _current_vid(ws_oid)
    svc = _service(client)
    p = svc.create_proposal(_req(ws_oid, base_vid))
    svc.approve_proposal(p["id"])
    first = svc.apply_proposal(p["id"])

    # 响应丢失后的安全重试：applied → 幂等返回既有结果（非报错）
    second = svc.apply_proposal(p["id"])
    assert second["status"] == "applied"
    assert second["applied_version_id"] == first["applied_version_id"]
    assert _ver_count(ws_oid) == 2, "重复 apply 产生了第二个版本"


# ---------------- 3a：draft 不可抢占（人工审批边界在 DB 层强制） ----------------


def test_draft_cannot_be_claimed(client):
    """项目负责人验收修正：draft 不在抢占条件中——即使绕过服务层前置检查，
    持久抢占层对 draft 的 claim rowcount 必须为 0（未经审批的 AI 内容不得写入）。"""
    ws_oid, _doc_id = _setup_workspace(client)
    base_vid = _current_vid(ws_oid)
    stub = _AiAppendStub()
    svc = _service(client, mutator=stub)
    p = svc.create_proposal(_req(ws_oid, base_vid))  # draft，未 approve

    # 直接执行抢占层同款条件更新（模拟服务层前置检查被遗漏）
    s = _db_session()
    try:
        rowcount = (s.query(AiProposal)
                    .filter(AiProposal.id == p["id"],
                            AiProposal.status.in_(("approved", "failed")))
                    .update({"status": "applying"}, synchronize_session=False))
        status = s.get(AiProposal, p["id"]).status
    finally:
        s.close()
    assert rowcount == 0, "draft 被 claim（审批边界失守）"
    assert status == "draft"

    # 服务层 apply(draft) → 拒绝且 mutator 未执行
    with pytest.raises(Exception) as ei:
        svc.apply_proposal(p["id"])
    assert ei.value.code == "PROPOSAL_INVALID_STATE"
    assert "尚未审批" in str(ei.value)
    assert stub.calls == 0
    assert _ver_count(ws_oid) == 1


# ---------------- 3b：applying 抢占（DB 条件更新），未抢到不得执行 mutator ----------------


def test_applying_blocks_concurrent_apply(client):
    ws_oid, _doc_id = _setup_workspace(client)
    base_vid = _current_vid(ws_oid)
    stub = _AiAppendStub()
    svc = _service(client, mutator=stub)
    p = svc.create_proposal(_req(ws_oid, base_vid))
    svc.approve_proposal(p["id"])

    # 模拟并发：另一请求已把提案抢到 applying
    s = _db_session()
    try:
        s.get(AiProposal, p["id"]).status = "applying"
        s.commit()
    finally:
        s.close()

    with pytest.raises(Exception) as ei:
        svc.apply_proposal(p["id"])
    assert ei.value.code == "PROPOSAL_APPLYING"
    assert stub.calls == 0, "抢占失败方执行了 mutator"
    assert _ver_count(ws_oid) == 1


# ---------------- 3c：failed 可重试（明确重试策略） ----------------


def test_failed_retry_policy(client):
    ws_oid, _doc_id = _setup_workspace(client)
    base_vid = _current_vid(ws_oid)
    bad = _AiAppendStub(fail=True)
    good = _AiAppendStub()
    svc = _service(client, mutator=bad)
    p = svc.create_proposal(_req(ws_oid, base_vid))
    svc.approve_proposal(p["id"])

    from app.services.ai_proposal import WorkspaceMutationError

    with pytest.raises(WorkspaceMutationError) as ei:
        svc.apply_proposal(p["id"])
    assert "AI 写入失败" in str(ei.value)
    assert _db_session().get(AiProposal, p["id"]).status == "failed"

    # 重试策略：failed 可再次 apply（换好引擎后成功）
    svc2 = _service(client, mutator=good)
    applied = svc2.apply_proposal(p["id"])
    assert applied["status"] == "applied" and applied["applied_version_id"]
    assert _ver_count(ws_oid) == 2


# ---------------- 4：rejected 不可应用 ----------------


def test_reject_then_apply_blocked(client):
    ws_oid, _doc_id = _setup_workspace(client)
    base_vid = _current_vid(ws_oid)
    svc = _service(client)
    p = svc.create_proposal(_req(ws_oid, base_vid))
    svc.reject_proposal(p["id"])
    assert svc._get(p["id"]).status == "rejected"

    with pytest.raises(Exception) as ei:
        svc.apply_proposal(p["id"])
    assert ei.value.code == "PROPOSAL_INVALID_STATE"
    assert _ver_count(ws_oid) == 1


# ---------------- 5：base 失效 → expired（409）且不可再应用 ----------------


def test_stale_base_expires_proposal(client):
    ws_oid, _doc_id = _setup_workspace(client)
    base_vid = _current_vid(ws_oid)
    svc = _service(client)
    p = svc.create_proposal(_req(ws_oid, base_vid))
    svc.approve_proposal(p["id"])

    # 工作稿在提案后前移（模拟另一操作已落版）
    s = _db_session()
    try:
        od = s.get(OfficeDocument, ws_oid)
        from app.services.office_document_svc import (
            OfficeDocumentService,
            OfficeDocumentTableAdapter,
        )

        OfficeDocumentService(session=s, adapter=OfficeDocumentTableAdapter(s)).create_version(
            od, _docx_bytes("其他操作产生的新版本"), source="onlyoffice")
        s.expire(od)
    finally:
        s.close()

    with pytest.raises(Exception) as ei:
        svc.apply_proposal(p["id"])
    assert ei.value.code == "WORKSPACE_VERSION_CONFLICT"
    assert _db_session().get(AiProposal, p["id"]).status == "expired"

    with pytest.raises(Exception) as ei:
        svc.apply_proposal(p["id"])
    assert ei.value.code == "PROPOSAL_INVALID_STATE"  # expired 不可应用


# ---------------- 6：mutator 失败 → failed + 零副作用 ----------------


def test_mutator_failure_zero_side_effects(client):
    ws_oid, _doc_id = _setup_workspace(client)
    base_vid = _current_vid(ws_oid)
    svc = _service(client, mutator=_AiAppendStub(fail=True))
    p = svc.create_proposal(_req(ws_oid, base_vid))
    svc.approve_proposal(p["id"])

    ver_before = _ver_count(ws_oid)
    current_before = _current_vid(ws_oid)

    with pytest.raises(Exception) as ei:
        svc.apply_proposal(p["id"])
    assert "AI 写入失败" in str(ei.value)

    s = _db_session()
    try:
        row = s.get(AiProposal, p["id"])
        derived = (s.query(WorkspaceDocument)
                   .filter(WorkspaceDocument.office_document_id == ws_oid).one())
    finally:
        s.close()
    assert row.status == "failed" and "AI 写入失败" in (row.error or "")
    assert _ver_count(ws_oid) == ver_before == 1
    assert _current_vid(ws_oid) == current_before
    assert derived.derived_from_version_id == current_before  # 派生未受影响


# ---------------- 7：不写 Source / Material（含 content_text） ----------------


def test_no_source_material_writes(client):
    ws_oid, _doc_id = _setup_workspace(client)
    base_vid = _current_vid(ws_oid)

    s = _db_session()
    try:
        before = _fingerprint(s)
    finally:
        s.close()

    r = _service(client).create_proposal(_req(ws_oid, base_vid))
    svc = _service(client)
    svc.approve_proposal(r["id"])
    svc.apply_proposal(r["id"])

    s = _db_session()
    try:
        assert _fingerprint(s) == before, "Source / Material 被写入"
    finally:
        s.close()


# ---------------- 8：活跃编辑 → 423，提案保持可重试状态 ----------------


def test_active_edit_keeps_proposal(client):
    ws_oid, _doc_id = _setup_workspace(client)
    base_vid = _current_vid(ws_oid)
    svc = _service(client, probe=lambda _oid: True)
    p = svc.create_proposal(_req(ws_oid, base_vid))
    svc.approve_proposal(p["id"])

    with pytest.raises(Exception) as ei:
        svc.apply_proposal(p["id"])
    assert ei.value.code == "WORKSPACE_EDITING"
    assert _db_session().get(AiProposal, p["id"]).status == "approved"  # 未被破坏，可重试


# ---------------- 9：operation / content 校验 ----------------


def test_invalid_operation_and_content(client):
    ws_oid, _doc_id = _setup_workspace(client)
    base_vid = _current_vid(ws_oid)
    svc = _service(client)

    with pytest.raises(Exception) as ei:
        svc.create_proposal(_req(ws_oid, base_vid, operation="replace_range"))
    assert ei.value.code == "PROPOSAL_INVALID_OPERATION"
    with pytest.raises(Exception) as ei:
        svc.create_proposal(_req(ws_oid, base_vid, text="   "))
    assert ei.value.code == "PROPOSAL_INVALID_CONTENT"
    assert _db_session().query(AiProposal).count() == 0
