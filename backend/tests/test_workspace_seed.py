"""工作稿底稿初始化（原型工作台画布=投标文件预览）。

- 新建 Workspace：项目已有投标文件 → v1 内容=投标文件（Office 版本链主路径 / stored_name 兜底）
- 既有原始骨架（仅 1 个 init 版本）→ 一次性补种；重复 ensure 幂等
- 已加工工作稿（≥2 版本）不触碰
- 无投标文件 → blank 骨架
"""

from __future__ import annotations

import io
import os

import pytest
from docx import Document as Docx
from fastapi.testclient import TestClient
from sqlalchemy import text
from sqlalchemy.orm import sessionmaker

from app.db import get_session, make_engine
from app.main import app


@pytest.fixture()
def client(tmp_path, monkeypatch):
    os.environ["ACCESS_PASSWORD"] = ""
    os.environ["SOURCE_MAINTENANCE_MODE"] = "0"
    os.environ["DATA_DIR"] = str(tmp_path / "data")
    os.environ["SCAN_INTERVAL_SECONDS"] = "3600"
    import app.config as cfg
    import app.db as appdb

    monkeypatch.setattr(cfg, "_SETTINGS", None)
    monkeypatch.setattr(appdb, "_Engine", None)
    monkeypatch.setattr(appdb, "_SessionLocal", None)
    engine = make_engine(cfg.get_settings().db_path)
    app.dependency_overrides[get_session] = _session_factory(engine)
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def _session_factory(engine):
    factory = sessionmaker(bind=engine, expire_on_commit=False)

    def _make():
        s = factory()
        try:
            yield s
        finally:
            s.close()

    return _make


def _db():
    import app.config as cfg
    import app.db as appdb

    appdb._Engine = None
    appdb._SessionLocal = None
    factory = sessionmaker(bind=make_engine(cfg.get_settings().db_path), expire_on_commit=False)
    return factory()


def _docx_bytes(body: str) -> bytes:
    buf = io.BytesIO()
    d = Docx()
    d.add_heading("底稿文档", level=0)
    for line in body.split("\n"):
        d.add_paragraph(line)
    d.save(buf)
    return buf.getvalue()


def _mk_bid_doc(pid: str, filename: str, body: str, with_office: bool) -> str:
    """直连模型插入投标文档档案（不经拆解队列，保证确定性）。"""
    import app.config as cfg
    from app.models import Document
    from app.services.office_document_svc import OfficeDocumentService, OfficeDocumentTableAdapter

    data = _docx_bytes(body)
    files_dir = cfg.get_settings().files_dir
    stored = f"seed_{filename}"
    (files_dir / stored).write_bytes(data)
    s = _db()
    try:
        office_document_id = None
        if with_office:
            svc = OfficeDocumentService(session=s, adapter=OfficeDocumentTableAdapter(s))
            from app.models import OfficeDocument

            od = OfficeDocument(document_type="source", title=filename, file_type="docx")
            s.add(od)
            s.flush()
            import hashlib

            svc.adapter.write_file(stored, data)
            v = svc.adapter.insert_version(
                od.id, 1, file_path=stored, file_hash=hashlib.sha256(data).hexdigest(),
                file_size=len(data), source="upload", editor_user=None, comment=None,
                restored_from_version_id=None,
            )
            svc.adapter.set_current_version(od.id, v.id)
            office_document_id = od.id
        doc = Document(filename=filename, stored_name=stored, source="upload", kind="bid",
                       file_type="docx", size=len(data), project_id=pid, status=2,
                       office_document_id=office_document_id)
        s.add(doc)
        s.commit()
        return doc.id
    finally:
        s.close()


def _ensure(client, pid: str) -> dict:
    r = client.post(f"/api/projects/{pid}/workspace")
    assert r.status_code == 200, r.text
    return r.json()


def _versions(oid: str) -> list:
    s = _db()
    try:
        return s.execute(
            text("SELECT version, source, comment FROM office_document_versions "
                 "WHERE office_document_id=:i ORDER BY version"), {"i": oid},
        ).fetchall()
    finally:
        s.close()


def _derived(pid: str) -> str | None:
    s = _db()
    try:
        return s.execute(
            text("SELECT derived_text FROM workspace_documents WHERE project_id=:i"), {"i": pid},
        ).scalar()
    finally:
        s.close()


MARK = "投标底稿标记-技术方案正文"


def test_new_workspace_seeded_from_bid_office_chain(client):
    pid = client.post("/api/projects", json={"name": "底稿项目A"}).json()["id"]
    _mk_bid_doc(pid, "投标A.docx", MARK, with_office=True)
    r = _ensure(client, pid)
    assert r["version"] == 1
    blocks = client.get(f"/api/office-documents/{r['office_document_id']}/reader").json()["blocks"]
    assert any(b["type"] == "para" and MARK in b.get("text", "") for b in blocks), blocks
    assert MARK in (_derived(pid) or ""), "底稿派生文本未同步"
    vs = _versions(r["office_document_id"])
    assert len(vs) == 1 and "投标A.docx" in (vs[0][2] or "")


def test_new_workspace_seeded_fallback_stored_name(client):
    pid = client.post("/api/projects", json={"name": "底稿项目B"}).json()["id"]
    _mk_bid_doc(pid, "投标B.docx", MARK, with_office=False)
    r = _ensure(client, pid)
    blocks = client.get(f"/api/office-documents/{r['office_document_id']}/reader").json()["blocks"]
    assert any(b["type"] == "para" and MARK in b.get("text", "") for b in blocks)


def test_blank_skeleton_backfilled_once(client):
    pid = client.post("/api/projects", json={"name": "底稿项目C"}).json()["id"]
    r1 = _ensure(client, pid)  # 无投标文件 → blank v1
    oid = r1["office_document_id"]
    assert len(_versions(oid)) == 1
    _mk_bid_doc(pid, "投标C.docx", MARK, with_office=False)
    r2 = _ensure(client, pid)  # 原始骨架 → 一次性补种
    assert r2["office_document_id"] == oid
    vs = _versions(oid)
    assert len(vs) == 2 and vs[-1][1] == "init" and "投标C.docx" in (vs[-1][2] or "")
    blocks = client.get(f"/api/office-documents/{oid}/reader").json()["blocks"]
    assert any(b["type"] == "para" and MARK in b.get("text", "") for b in blocks)
    assert MARK in (_derived(pid) or "")
    r3 = _ensure(client, pid)  # 幂等：不重复补种
    assert r3["version"] == 2 and len(_versions(oid)) == 2


def test_processed_workspace_untouched(client):
    import hashlib

    from app.models import OfficeDocument
    from app.services.office_document_svc import OfficeDocumentService, OfficeDocumentTableAdapter

    pid = client.post("/api/projects", json={"name": "底稿项目D"}).json()["id"]
    r1 = _ensure(client, pid)
    oid = r1["office_document_id"]
    # 模拟用户已加工：追加 material_insert v2
    s = _db()
    try:
        svc = OfficeDocumentService(session=s, adapter=OfficeDocumentTableAdapter(s))
        data = _docx_bytes("已插入素材内容")
        stored = f"workspace_{oid[:8]}_2.docx"
        svc.adapter.write_file(stored, data)
        v2 = svc.adapter.insert_version(
            oid, 2, file_path=stored, file_hash=hashlib.sha256(data).hexdigest(),
            file_size=len(data), source="material_insert", editor_user=None,
            comment="素材插入", restored_from_version_id=None,
        )
        svc.adapter.set_current_version(oid, v2.id)
        s.commit()
        cur_before = s.get(OfficeDocument, oid).current_version_id
    finally:
        s.close()
    _mk_bid_doc(pid, "投标D.docx", MARK, with_office=False)
    r2 = _ensure(client, pid)  # 已加工 → 不触碰
    assert r2["office_document_id"] == oid and r2["version"] == 2
    s = _db()
    try:
        assert s.get(OfficeDocument, oid).current_version_id == cur_before
        assert len(_versions(oid)) == 2
    finally:
        s.close()


def test_no_bid_blank_skeleton(client):
    pid = client.post("/api/projects", json={"name": "底稿项目E"}).json()["id"]
    r = _ensure(client, pid)
    blocks = client.get(f"/api/office-documents/{r['office_document_id']}/reader").json()["blocks"]
    assert not any(MARK in b.get("text", "") for b in blocks)
    assert len(_versions(r["office_document_id"])) == 1
