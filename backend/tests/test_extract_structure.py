"""拆解结构回归（原样复制制）：正文不丢、冒号不拆列、真实表格保结构、编辑与重拆幂等。

对应外部审查 Case 1~6：
1 无冒号普通正文必须入库；2 有冒号正文 cells_json 保持单列；
3 真实三列表格保持 2 行 3 列；4 段落+表格+段落顺序保持；
5 编辑保存 round-trip 后仍为段落；6 重复重新拆解结果一致。
"""

from __future__ import annotations

import hashlib
import io
import json
import os
import time

import pytest
from docx import Document as Docx
from fastapi.testclient import TestClient

from app.db import get_session, make_engine
from app.main import app

TENDER_STRUCT = [
    "投标人须知前附表",
    "1.1.2 招标人：见投标人须知前附表。",
    "本项目位于北京市朝阳区。",
]

TABLE_CELLS = [
    ["9", "是否采用电子招标投标", "否，具体要求见商务文件。"],
    ["10", "需要补充的其他内容", "投标人应当承诺满足全部技术条款。"],
]

TAIL_PARA = "表格后的正文"


@pytest.fixture()
def client(tmp_path, monkeypatch):
    os.environ["ACCESS_PASSWORD"] = ""
    os.environ["DATA_DIR"] = str(tmp_path / "data")
    os.environ["SCAN_INTERVAL_SECONDS"] = "3600"
    import app.config as cfg
    import app.db as appdb

    monkeypatch.setattr(cfg, "_SETTINGS", None)
    monkeypatch.setattr(appdb, "_Engine", None)
    monkeypatch.setattr(appdb, "_SessionLocal", None)
    engine = make_engine(cfg.get_settings().db_path)
    app.dependency_overrides[get_session] = _session_factory(engine)
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def _session_factory(engine):
    from sqlalchemy.orm import sessionmaker

    factory = sessionmaker(bind=engine, expire_on_commit=False)

    def _make():
        s = factory()
        try:
            yield s
        finally:
            s.close()

    return _make


def _struct_docx() -> bytes:
    buf = io.BytesIO()
    d = Docx()
    for line in TENDER_STRUCT:
        d.add_paragraph(line)
    tb = d.add_table(rows=len(TABLE_CELLS), cols=3)
    for ri, row in enumerate(TABLE_CELLS):
        for ci, val in enumerate(row):
            tb.cell(ri, ci).text = val
    d.add_paragraph(TAIL_PARA)
    d.save(buf)
    return buf.getvalue()


def _upload_struct_doc(client) -> str:
    pid = client.post("/api/projects", json={"name": "结构回归项目"}).json()["id"]
    r = client.post(
        "/api/documents/upload",
        files=[("files", ("结构回归.docx", _struct_docx(), "application/octet-stream"))],
        data={"kind": "tender", "project_id": pid, "is_secret": "false"},
    )
    assert r.status_code == 200
    doc_id = r.json()["created"][0]["id"]
    deadline = time.time() + 30
    while time.time() < deadline:
        docs = client.get("/api/documents").json()["items"]
        d = next(x for x in docs if x["id"] == doc_id)
        if d["status"] == 2:
            break
        time.sleep(0.2)
    assert d["status"] == 2, d
    return doc_id


def _rows(client, doc_id: str) -> list[list[str]]:
    ex = client.get(f"/api/documents/{doc_id}/extract").json()
    return [row for g in ex["groups"] for it in g["items"] for row in it["rows"]]


def test_case1_colonless_paragraph_persisted(client):
    doc_id = _upload_struct_doc(client)
    rows = _rows(client, doc_id)
    assert ["本项目位于北京市朝阳区。"] in rows, f"无冒号正文丢失：{rows}"


def test_case2_colon_paragraph_stays_single_column(client):
    doc_id = _upload_struct_doc(client)
    rows = _rows(client, doc_id)
    assert ["1.1.2 招标人：见投标人须知前附表。"] in rows, f"冒号正文被拆列：{rows}"
    assert ["1.1.2 招标人", "见投标人须知前附表。"] not in rows


def test_case3_real_table_kept(client):
    doc_id = _upload_struct_doc(client)
    rows = _rows(client, doc_id)
    for row in TABLE_CELLS:
        assert row in rows, f"表格行失真：{rows}"
    multi = [r for r in rows if len(r) == 3]
    assert len(multi) >= 2, "三列表格行数不足"


def test_case4_para_table_para_order(client):
    doc_id = _upload_struct_doc(client)
    rows = _rows(client, doc_id)
    i_anchor = rows.index(["投标人须知前附表"])
    i_colon = rows.index(["1.1.2 招标人：见投标人须知前附表。"])
    i_plain = rows.index(["本项目位于北京市朝阳区。"])
    i_tbl = rows.index(TABLE_CELLS[0])
    i_tail = rows.index([TAIL_PARA])
    assert i_anchor < i_colon < i_plain < i_tbl < i_tail, f"顺序错乱：{rows}"


def test_case5_edit_roundtrip_stays_paragraph(client):
    doc_id = _upload_struct_doc(client)
    ex = client.get(f"/api/documents/{doc_id}/extract").json()
    item = ex["groups"][0]["items"][0]
    rows = [list(r) for r in item["rows"]]
    new_rows = [
        ["1.1.2 招标人：中国测试二有限公司。"] if r == ["1.1.2 招标人：见投标人须知前附表。"] else r
        for r in rows
    ]
    r = client.put(f"/api/documents/{doc_id}/extract/{item['id']}", json={"rows": new_rows})
    assert r.status_code == 200
    rows2 = _rows(client, doc_id)
    assert ["1.1.2 招标人：中国测试二有限公司。"] in rows2, f"编辑后正文丢失或被拆列：{rows2}"


def test_case6_reextract_idempotent(client):
    doc_id = _upload_struct_doc(client)

    def snap() -> tuple[int, str]:
        rows = _rows(client, doc_id)
        h = hashlib.sha1(json.dumps(rows, ensure_ascii=False).encode("utf-8")).hexdigest()
        return len(rows), h

    s1 = snap()
    for _ in range(2):
        r = client.post(f"/api/documents/{doc_id}/extract")
        assert r.status_code == 200
    s2 = snap()
    assert s1 == s2, f"重复重拆不一致：{s1} -> {s2}"
