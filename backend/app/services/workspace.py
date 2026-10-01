"""C-1：Workspace OfficeDocument 底座。

- 每个 Project 最多一份 Workspace（OfficeDocument type=workspace）
- Workspace DOCX 是投标文件工作正文唯一真源；本模块只做：
  初始化（blank/template DOCX v1）与派生数据（WorkspaceDocumentProcessor）
- 版本/历史/恢复/callback/key/token 全部复用 OfficeDocumentService，不新建专用服务
"""

from __future__ import annotations

import io
import logging
import threading
from typing import Any

from sqlalchemy.orm import Session

logger = logging.getLogger(__name__)

# ensure 并发双防：页面加载可并发触发两次 ensure，无锁会创建两个 Workspace（实机孤儿案例）
_ensure_lock = threading.Lock()


def blank_workspace_bytes(title: str) -> bytes:
    """初始化用 blank 工作稿（合法可编辑 DOCX；正式模板后续替换）。"""
    from docx import Document as Docx

    d = Docx()
    d.add_heading(title, level=0)
    d.add_paragraph("")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _bid_seed_bytes(session: Session, project: Any) -> tuple[bytes, str] | None:
    """项目最新投标文件 → 工作稿底稿内容；无投标文件或均不可读则 None。

    文件读取按 B-1c 真源顺序：OfficeDocument 版本链 → stored_name（deprecated 兜底）→ origin_path。
    """
    from pathlib import Path

    from ..config import get_settings
    from ..models import Document, OfficeDocument, OfficeDocumentVersion
    from .office_document_svc import OfficeDocumentService, OfficeDocumentTableAdapter

    rows = (session.query(Document)
            .filter(Document.project_id == project.id, Document.kind == "bid")
            .order_by(Document.created_at.desc()).all())
    files_dir = get_settings().files_dir
    svc = None
    for doc in rows:
        if doc.office_document_id and svc is None:
            svc = OfficeDocumentService(session=session, adapter=OfficeDocumentTableAdapter(session))
        if doc.office_document_id:
            od = session.get(OfficeDocument, doc.office_document_id)
            cur = (session.get(OfficeDocumentVersion, od.current_version_id)
                   if od is not None and od.current_version_id else None)
            if cur is not None:
                try:
                    return svc.adapter.read_file(cur.file_path), doc.filename
                except Exception:
                    pass
        cands: list[Path] = []
        if doc.stored_name:
            cands.append(files_dir / doc.stored_name)
        if doc.origin_path:
            cands.append(Path(doc.origin_path))
        for cand in cands:
            try:
                if not cand.is_file():
                    continue
                data = cand.read_bytes()
            except OSError:
                continue
            if 0 < len(data) <= 200 * 1024 * 1024:
                return data, doc.filename
    return None


def _sync_workspace_derived(session: Session, od: Any, v: Any) -> None:
    """底稿非 blank，同步重算派生文本（失败不阻塞初始化，日志登记人工排查）。"""
    try:
        process_workspace_document(session, od, v)
    except Exception:
        logger.exception("workspace seed derived sync failed: od=%s v=%s", od.id, v.version)
        session.rollback()


def _backfill_bid_seed(session: Session, od: Any, seed: tuple[bytes, str] | None) -> None:
    """原始骨架工作稿（仅 1 个 init 版本、从未加工）+ 项目已有投标文件 → 一次性补种底稿。

    已加工过（素材插入/AI 写入等 ≥2 版本）不触碰；同内容幂等跳过。
    """
    import hashlib

    from ..models import OfficeDocumentVersion
    from .audit import audit
    from .office_document_svc import OfficeDocumentService, OfficeDocumentTableAdapter

    if seed is None:
        return
    versions = (session.query(OfficeDocumentVersion)
                .filter(OfficeDocumentVersion.office_document_id == od.id)
                .all())
    if len(versions) != 1 or versions[0].source != "init":
        return
    v1 = versions[0]
    data, src_name = seed
    file_hash = hashlib.sha256(data).hexdigest()
    if v1.file_hash == file_hash:
        return
    svc = OfficeDocumentService(session=session, adapter=OfficeDocumentTableAdapter(session))
    stored = f"workspace_{od.id[:8]}_{v1.version + 1}.docx"
    svc.adapter.write_file(stored, data)
    v2 = svc.adapter.insert_version(
        od.id, v1.version + 1, file_path=stored, file_hash=file_hash,
        file_size=len(data), source="init", editor_user=None,
        comment=f"以投标文件「{src_name}」初始化底稿", restored_from_version_id=None,
    )
    svc.adapter.set_current_version(od.id, v2.id)
    session.commit()
    _sync_workspace_derived(session, od, v2)
    audit("workspace_seeded", detail=f"office_document {od.id} v{v2.version} · {src_name}")


def ensure_workspace(session: Session, project_id: str) -> Any:
    """幂等：返回项目 Workspace 的 OfficeDocument（不存在则创建）。

    v1 内容优先取项目最新投标文件为底稿（原型工作台画布=投标文件预览）；
    无投标文件时退回 blank 骨架。既有原始骨架（从未加工）在获得投标文件后一次性补种。
    禁止重复创建两个 Workspace（以 project.workspace_office_document_id 为准）。
    """
    from ..models import OfficeDocument, Project, WorkspaceDocument
    from .office_document_svc import OfficeDocumentService, OfficeDocumentTableAdapter

    project = session.get(Project, project_id)
    if project is None:
        raise ValueError("项目不存在")
    title = f"{project.name} 投标工作稿"
    with _ensure_lock:
        return _ensure_locked(session, project, title)


def _ensure_locked(session: Session, project: Any, title: str) -> Any:
    from ..models import OfficeDocument, WorkspaceDocument
    from .office_document_svc import OfficeDocumentService, OfficeDocumentTableAdapter

    seed = _bid_seed_bytes(session, project)
    if project.workspace_office_document_id:
        od = session.get(OfficeDocument, project.workspace_office_document_id)
        if od is not None:
            _backfill_bid_seed(session, od, seed)
            return od
    if seed is not None:
        data, src_name = seed
        comment = f"以投标文件「{src_name}」初始化底稿"
    else:
        data, comment = blank_workspace_bytes(title), "工作稿初始化"
    od = OfficeDocument(document_type="workspace", title=title, file_type="docx")
    session.add(od)
    session.flush()
    svc = OfficeDocumentService(session=session, adapter=OfficeDocumentTableAdapter(session))
    stored = f"workspace_{od.id[:8]}.docx"
    svc.adapter.write_file(stored, data)
    import hashlib

    v = svc.adapter.insert_version(
        od.id, 1, file_path=stored, file_hash=hashlib.sha256(data).hexdigest(),
        file_size=len(data), source="init", editor_user=None, comment=comment,
        restored_from_version_id=None,
    )
    svc.adapter.set_current_version(od.id, v.id)
    project.workspace_office_document_id = od.id
    session.add(WorkspaceDocument(
        project_id=project.id, office_document_id=od.id,
        derived_from_version_id=v.id,
    ))
    session.commit()
    if seed is not None:
        _sync_workspace_derived(session, od, v)
    logger.info("workspace created: project=%s od=%s seeded=%s",
                project.id, od.id, bool(seed))
    return od


def process_workspace_document(session: Any, od: Any, version: Any) -> None:
    """Workspace 派生数据：parse DOCX → derived_text（正文真源仍是 OfficeDocumentVersion）。"""
    from ..config import get_settings
    from ..models import WorkspaceDocument
    from .ingest.parser import parse_docx

    row = (
        session.query(WorkspaceDocument)
        .filter(WorkspaceDocument.office_document_id == od.id)
        .first()
    )
    if row is None:
        logger.warning("workspace row for office_document %s not found", od.id)
        return
    data = (get_settings().files_dir / version.file_path).read_bytes()
    parsed = parse_docx(data)
    lines: list[str] = []
    for tb in parsed.doc_tables:
        model = tb.model
        for r in (model["rows"] if model else []):
            cells = [c["text"].strip() for c in r["cells"] if not c.get("merged")]
            cells = [c for c in cells if c]
            if cells:
                lines.append(" | ".join(cells))
    for ln in parsed.text.split("\n"):
        s = ln.strip()
        if s:
            lines.append(s)
    row.derived_text = "\n".join(lines) or None
    row.derived_from_version_id = version.id
    session.commit()
    logger.info("workspace derived text updated: od=%s v=%s lines=%d", od.id, version.version, len(lines))


class WorkspaceDocumentProcessor:
    """post-save processor：document_type = workspace。"""

    document_type = "workspace"

    def process(self, od: Any, version: Any) -> None:
        from sqlalchemy.orm import sessionmaker

        from ..db import get_engine

        factory = sessionmaker(bind=get_engine(), expire_on_commit=False)
        ses = factory()
        try:
            process_workspace_document(ses, od, version)
        finally:
            ses.close()
