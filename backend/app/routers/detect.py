"""F3.1/F3.3 废标检测路由：运行检测 / 风险清单 / 已处理闭环 / 检测报告 Excel。"""

from __future__ import annotations

import io
import json
from typing import Any

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy.orm import Session

from ..db import get_session
from ..models import Document, RiskItem
from ..services.audit import audit
from ..services.detect import BUILTIN_RULES, detect_project, ensure_rules, risk_index_of

router = APIRouter(prefix="/api/detect", tags=["F3 废标检测"])


@router.get("/rules")
def list_rules(session: Session = Depends(get_session)) -> dict[str, Any]:
    """检测规则清单（库为准、内置顺序），含启用状态。"""
    from ..models import DetectRule

    ensure_rules(session)
    rows = {r.id: r for r in session.query(DetectRule).all()}
    rules = [
        {
            "id": r["id"], "name": r["name"], "sev": r["sev"], "type": r["type"],
            "desc": rows[r["id"]].desc if r["id"] in rows else r["desc"],
            "enabled": rows[r["id"]].enabled if r["id"] in rows else True,
        }
        for r in BUILTIN_RULES
    ]
    return {"rules": rules}


class RulePatchIn(BaseModel):
    enabled: bool | None = None
    desc: str | None = None


@router.patch("/rules/{rule_id}")
def patch_rule(rule_id: str, body: RulePatchIn, session: Session = Depends(get_session)) -> dict[str, Any]:
    """规则维护：启停/改说明，每次变更写入留痕表。"""
    from ..models import DetectRule, DetectRuleVersion

    row = session.get(DetectRule, rule_id)
    if row is None:
        raise HTTPException(404, "规则不存在")
    changes: list[str] = []
    if body.enabled is not None and body.enabled != row.enabled:
        row.enabled = body.enabled
        changes.append("启用" if body.enabled else "停用")
    if body.desc is not None and body.desc.strip() and body.desc.strip() != row.desc:
        row.desc = body.desc.strip()
        changes.append("说明修改")
    if not changes:
        return {"ok": True, "changed": False}
    session.add(DetectRuleVersion(
        rule_id=rule_id,
        note="；".join(changes),
        snapshot_json=json.dumps(
            {"name": row.name, "sev": row.sev, "type": row.type, "desc": row.desc, "enabled": row.enabled},
            ensure_ascii=False,
        ),
    ))
    session.commit()
    audit("detect_rule_updated", detail=f"{rule_id} · {'；'.join(changes)}")
    return {"ok": True, "changed": True}


@router.get("/rules/history")
def rule_history(limit: int = 30, session: Session = Depends(get_session)) -> dict[str, Any]:
    """规则变更留痕。"""
    from ..models import DetectRuleVersion

    rows = session.query(DetectRuleVersion).order_by(DetectRuleVersion.created_at.desc()).limit(limit).all()
    names = {r["id"]: r["name"] for r in BUILTIN_RULES}
    return {"items": [
        {
            "rule_id": v.rule_id, "rule_name": names.get(v.rule_id, v.rule_id),
            "note": v.note,
            "created_at": v.created_at.isoformat() if v.created_at else None,
        }
        for v in rows
    ]}


def _doc_names(session: Session) -> dict[str, str]:
    return {d.id: d.filename for d in session.query(Document).all()}


def _risk_dict(r: RiskItem, names: dict[str, str]) -> dict[str, Any]:
    return {
        "id": r.id, "sev": r.sev, "type": r.type, "title": r.title,
        "meta": r.meta, "req": r.req, "quote": r.quote, "advice": r.advice,
        "fix": r.fix, "page_no": r.page_no, "done": bool(r.done),
        "doc_id": r.doc_id, "doc_name": names.get(r.doc_id or ""),
        "created_at": r.created_at.isoformat() if r.created_at else None,
    }


@router.post("/{project_id}")
def run_detect(project_id: str, session: Session = Depends(get_session)) -> dict[str, Any]:
    try:
        result = detect_project(session, project_id)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    items = session.query(RiskItem).filter(RiskItem.project_id == project_id).order_by(RiskItem.created_at).all()
    names = _doc_names(session)
    return {
        "risk_index": risk_index_of(items),
        "detected_at": result.detected_at,
        "score_rows": result.score_rows,
        "items": [_risk_dict(r, names) for r in items],
    }


@router.get("/{project_id}/risks")
def list_risks(project_id: str, session: Session = Depends(get_session)) -> dict[str, Any]:
    from ..models import AuditLog

    items = session.query(RiskItem).filter(RiskItem.project_id == project_id).order_by(RiskItem.created_at).all()
    names = _doc_names(session)
    # has_result：有风险行 OR 有检测运行审计（无风险时 RiskItem 为空也要显示"未发现风险"）
    last_run = (
        session.query(AuditLog.ts)
        .filter(AuditLog.action == "detect_run", AuditLog.detail.like(f"%{project_id}%"))
        .order_by(AuditLog.ts.desc())
        .first()
    )
    has_result = bool(items) or last_run is not None
    latest = str(last_run[0]) if last_run else (items[-1].created_at.isoformat() if items else None)
    return {
        "risk_index": risk_index_of(items),
        "detected_at": latest,
        "has_result": has_result,
        "items": [_risk_dict(r, names) for r in items],
    }


class DoneIn(BaseModel):
    done: bool


@router.patch("/risks/{risk_id}/done")
def mark_done(risk_id: str, body: DoneIn, session: Session = Depends(get_session)) -> dict[str, Any]:
    r = session.get(RiskItem, risk_id)
    if r is None:
        raise HTTPException(404, "风险项不存在")
    r.done = 1 if body.done else 0
    session.commit()
    audit("risk_marked_done" if body.done else "risk_reopened", detail=r.title[:80])
    items = session.query(RiskItem).filter(RiskItem.project_id == r.project_id).all()
    return {"ok": True, "risk_index": risk_index_of(items)}


@router.get("/{project_id}/report.xlsx")
def report(project_id: str, session: Session = Depends(get_session)) -> Response:
    """检测报告 Excel（F3.3）。"""
    items = session.query(RiskItem).filter(RiskItem.project_id == project_id).order_by(RiskItem.created_at).all()
    if not items:
        raise HTTPException(404, "尚未运行检测")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "废标检测报告"
    ws.append(["废标风险项", "级别", "类型", "出处", "依据（招标条款）", "投标原文引用", "修改建议", "状态"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="EDF3FD")
    sev_label = {"fatal": "致命", "risk": "风险", "hint": "提示"}
    for r in items:
        ws.append([
            r.title, sev_label.get(r.sev, r.sev), r.type, r.meta or "",
            r.req or "", r.quote or "", r.advice or "", "已处理" if r.done else "未处理",
        ])
    for col, w in zip("ABCDEFGH", (34, 8, 8, 40, 40, 40, 44, 8), strict=False):
        ws.column_dimensions[col].width = w
    ws.append([])
    ws.append([f"废标风险指数：{risk_index_of(items)} / 100", f"共 {len(items)} 项 · 未处理 {sum(1 for x in items if not x.done)} 项"])
    buf = io.BytesIO()
    wb.save(buf)
    audit("report_exported", detail=f"project {project_id} · {len(items)} 项")
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=detect_report.xlsx"},
    )
