"""P3 文档库路由：列表/详情/拆解/上传/绑定。"""

from __future__ import annotations

import json
import re
import uuid
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel
from sqlalchemy.orm import Session

from ..config import get_settings
from ..db import get_session
from ..models import Document, ExtractItem, ExtractRow, Project
from ..services.audit import audit
from ..services.ingest.pipeline import process_document

router = APIRouter(prefix="/api/documents", tags=["F1 文档库"])

_MAX_SIZE = 200 * 1024 * 1024

_CHAPTER_TITLE = re.compile(r"^第[0-9一二三四五六七八九十百零]+[章节部分篇]")


class UploadIn(BaseModel):
    doc_ids: list[str] = []


@router.get("")
def list_docs(
    kind: str | None = None,
    project_id: str | None = None,
    q: str | None = None,
    status: int | None = None,
    file_type: str | None = None,
    session: Session = Depends(get_session),
) -> dict[str, Any]:
    from ..models import Document

    query = session.query(Document)
    if kind and kind in ("tender", "bid", "notice"):
        query = query.filter(Document.kind == kind)
    if project_id == "none":
        query = query.filter(Document.project_id.is_(None))
    elif project_id:
        query = query.filter(Document.project_id == project_id)
    if status is not None:
        query = query.filter(Document.status == status)
    if file_type == "img":
        query = query.filter(Document.file_type.in_(("jpg", "jpeg", "png", "scan")))
    elif file_type:
        query = query.filter(Document.file_type == file_type)
    if q:
        query = query.filter(Document.filename.contains(q))
    rows = query.order_by(Document.created_at.desc()).all()
    return {
        "total": len(rows),
        "items": [
            {
                "id": d.id, "filename": d.filename, "kind": d.kind,
                "file_type": d.file_type, "size": d.size,
                "status": d.status, "project_id": d.project_id,
                "is_secret": d.is_secret, "error_msg": d.error_msg,
                "ocr_conf": d.ocr_conf, "needs_review": d.needs_review,
                "created_at": d.created_at.isoformat() if d.created_at else None,
            }
            for d in rows
        ],
    }


@router.get("/{doc_id}/extract")
def get_extract(doc_id: str, session: Session = Depends(get_session)) -> dict[str, Any]:
    """返回拆解结构（按组聚合）。"""
    from ..models import Document

    doc = session.get(Document, doc_id)
    if doc is None:
        raise HTTPException(404, "文档不存在")
    if doc.status != 2:
        return {"groups": [], "status": doc.status}
    items = (
        session.query(ExtractItem)
        .filter(ExtractItem.doc_id == doc_id)
        .order_by(ExtractItem.seq)
        .all()
    )
    groups: dict[str, list[dict[str, Any]]] = {}
    for item in items:
        rows = (
            session.query(ExtractRow)
            .filter(ExtractRow.item_id == item.id)
            .order_by(ExtractRow.seq)
            .all()
        )
        out_rows: list[Any] = []  # 段落行=list[str]；V2 表格行=dict
        for r in rows:
            if r.cells_json:
                try:
                    cells = json.loads(r.cells_json)
                    if isinstance(cells, dict) and cells.get("type") == "table":
                        out_rows.append(cells)  # V2 表格行对象原样透传
                        continue
                    if isinstance(cells, list) and cells:
                        out_rows.append([str(c) for c in cells])
                        continue
                except ValueError:
                    pass
            out_rows.append([r.field, r.value or ""])
        groups.setdefault(item.group_name, []).append({
            "id": item.id, "name": item.name, "sum": item.sum,
            "rows": out_rows,
        })
    return {
        "status": doc.status,
        "groups": [{"name": g, "items": items_} for g, items_ in groups.items()],
    }


@router.post("/{doc_id}/extract")
def run_extract(doc_id: str, session: Session = Depends(get_session)) -> dict[str, Any]:
    doc = session.get(Document, doc_id)
    if doc is None:
        raise HTTPException(404, "文档不存在")
    if doc.status == 1:
        raise HTTPException(409, "拆解进行中")
    result = process_document(session, doc)
    return {"status": result.status, "extracted": result.extracted, "materials": result.materials}


class BindIn(BaseModel):
    project_id: str | None = None  # None = 解绑
    new_project_name: str | None = None
    new_project_industry: str | None = None


@router.post("/{doc_id}/bind")
def bind_project(doc_id: str, body: BindIn, session: Session = Depends(get_session)) -> dict[str, Any]:
    from ..models import BUILTIN_INDUSTRIES

    doc = session.get(Document, doc_id)
    if doc is None:
        raise HTTPException(404, "文档不存在")
    if body.new_project_name:
        from .projects import _unique_code

        np = Project(
            code=_unique_code(session, "NEW-" + body.new_project_name[:8]),
            name=body.new_project_name.strip(),
            industry=body.new_project_industry or BUILTIN_INDUSTRIES[0],
        )
        session.add(np)
        session.flush()
        doc.project_id = np.id
        audit("doc_bound_new_project", doc_id=doc.id, detail=np.name)
    elif body.project_id:
        doc.project_id = body.project_id
        audit("doc_bound", doc_id=doc.id, detail=body.project_id)
    else:
        doc.project_id = None
        audit("doc_unbound", doc_id=doc.id)
    session.commit()
    return {"ok": True, "project_id": doc.project_id}


class RowsIn(BaseModel):
    rows: list[Any]  # 段落行=字符串数组；V2 表格行=对象（type=table）


@router.put("/{doc_id}/extract/{item_id}")
def update_extract_item(doc_id: str, item_id: str, body: RowsIn, session: Session = Depends(get_session)) -> dict[str, Any]:
    """页内编辑保存：全量替换结构项的字段行。"""
    from ..models import ExtractRow

    item = session.get(ExtractItem, item_id)
    if item is None or item.doc_id != doc_id:
        raise HTTPException(404, "结构项不存在")
    session.query(ExtractRow).filter(ExtractRow.item_id == item.id).delete()
    for ri, row in enumerate(body.rows):
        if isinstance(row, dict):
            # V2 表格行：按 cell.id 回填后的对象原样存储（只允许 text 变化，结构不漂移）
            texts = [
                str(c.get("text", "")).strip()
                for c in row.get("cells", []) if not c.get("merged")
            ]
            texts = [t for t in texts if t]
            fname = " ".join(texts[:-1]) if len(texts) > 1 else (texts[0] if texts else "表格内容")
            value = texts[-1] if texts else ""
            session.add(ExtractRow(
                item_id=item.id, field=fname[:120] or "表格内容", value=value[:2000], seq=ri,
                cells_json=json.dumps(row, ensure_ascii=False),
            ))
            continue
        if not isinstance(row, list) or not row:
            continue
        cells = [str(c).strip() for c in row]
        field = cells[0][:120]
        if not field:
            continue
        value = cells[-1][:2000] if len(cells) > 1 else ""
        session.add(ExtractRow(
            item_id=item.id, field=field, value=value, seq=ri,
            cells_json=json.dumps([str(c) for c in row], ensure_ascii=False),
        ))
    item.sum = f"{len(body.rows)} 字段" if body.rows else item.sum
    session.commit()
    audit("extract_item_edited", doc_id=doc_id, detail=item.name[:60])
    return {"ok": True, "rows": len(body.rows)}


@router.get("/{doc_id}/content")
def doc_content(doc_id: str, session: Session = Depends(get_session)) -> dict[str, Any]:
    """正文预览：以 content_text 为准（编辑后的正文），无则回退切片拼接。"""
    from ..models import Chunk

    doc = session.get(Document, doc_id)
    if doc is None:
        raise HTTPException(404, "文档不存在")
    text = (doc.content_text or "").strip()
    if not text:
        rows = (
            session.query(Chunk)
            .filter(Chunk.document_id == doc_id)
            .order_by(Chunk.chunk_index)
            .all()
        )
        text = "\n".join(r.content for r in rows)
    return {"filename": doc.filename, "text": text[:60000]}


class ContentIn(BaseModel):
    text: str


@router.put("/{doc_id}/content")
def save_content(doc_id: str, body: ContentIn, session: Session = Depends(get_session)) -> dict[str, Any]:
    """工作台正文编辑保存：更新正文并重建检索索引（原件文件不动）。"""
    from ..services.ingest.pipeline import _write_index, purge_document_index

    doc = session.get(Document, doc_id)
    if doc is None:
        raise HTTPException(404, "文档不存在")
    doc.content_text = body.text
    purge_document_index(session, doc_id)
    _write_index(session, doc_id, [body.text])
    session.commit()
    audit("content_saved", doc_id=doc.id, detail=f"{len(body.text)} 字")
    return {"ok": True, "chars": len(body.text)}


@router.get("/{doc_id}/content.docx")
def content_docx(doc_id: str, session: Session = Depends(get_session)) -> Response:
    """成品稿导出：编辑后正文用 python-docx 生成 Word 下载（原件保留不动）。"""
    from io import BytesIO
    from urllib.parse import quote

    from docx import Document as Docx
    from docx.shared import Pt

    from ..models import Chunk

    doc = session.get(Document, doc_id)
    if doc is None:
        raise HTTPException(404, "文档不存在")
    text = (doc.content_text or "").strip()
    if not text:
        rows = (
            session.query(Chunk)
            .filter(Chunk.document_id == doc_id)
            .order_by(Chunk.chunk_index)
            .all()
        )
        text = "\n".join(r.content for r in rows)
    d = Docx()
    for line in text.splitlines():
        ln = line.strip()
        if not ln:
            continue
        p = d.add_paragraph()
        if _CHAPTER_TITLE.match(ln):
            run = p.add_run(ln)
            run.bold = True
            run.font.size = Pt(14)
        else:
            run = p.add_run(ln)
            run.font.size = Pt(12)
    buf = BytesIO()
    d.save(buf)
    fname = f"编辑稿_{doc.filename}" if doc.filename.endswith(".docx") else f"编辑稿_{doc.filename}.docx"
    audit("doc_exported", doc_id=doc.id, detail=fname)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"},
    )


@router.get("/{doc_id}/download")
def download_doc(doc_id: str, session: Session = Depends(get_session)) -> FileResponse:
    """成品导出：下载文档原件（Word/PDF 原样，F3.3）。"""
    doc = session.get(Document, doc_id)
    if doc is None:
        raise HTTPException(404, "文档不存在")
    settings = get_settings()
    if doc.stored_name:
        path = settings.files_dir / doc.stored_name
        if path.is_file():
            audit("doc_exported", doc_id=doc.id, detail=doc.filename)
            return FileResponse(path, filename=doc.filename)
    if doc.origin_path and Path(doc.origin_path).is_file():
        audit("doc_exported", doc_id=doc.id, detail=doc.filename)
        return FileResponse(Path(doc.origin_path), filename=doc.filename)
    raise HTTPException(404, "原件文件不存在")


@router.post("/upload")
async def upload(
    files: list[UploadFile] = File(...),
    kind: str = Form("tender"),
    project_id: str | None = Form(None),
    session: Session = Depends(get_session),
) -> dict[str, Any]:
    if kind not in ("tender", "bid", "notice"):
        raise HTTPException(400, "kind 须为 tender/bid/notice")
    allowed = {".docx", ".pdf", ".txt", ".xlsx", ".xls", ".csv", ".jpg", ".jpeg", ".png", ".zip"}
    for f in files:
        ext = Path(f.filename or "").suffix.lower()
        if ext not in allowed:
            raise HTTPException(415, f"不支持的文件类型：{f.filename}（允许：docx/pdf/txt/xlsx/xls/csv/jpg/png/zip）")
    settings = get_settings()
    created: list[dict[str, Any]] = []
    for f in files:
        data = await f.read()
        if len(data) > 200 * 1024 * 1024:
            raise HTTPException(413, f"{f.filename} 超过 200MB 限制")
        stored = uuid.uuid4().hex + Path(f.filename or "unnamed").suffix.lower()
        (settings.files_dir / stored).write_bytes(data)
        from ..models import Document as Doc

        doc = Doc(
            filename=(f.filename or "unnamed")[:500],
            stored_name=stored,
            source="upload",
            kind=kind,
            file_type=detect_file_type(f.filename or ""),
            size=len(data),
            project_id=project_id or None,
            status=0,
        )
        session.add(doc)
        session.commit()
        created.append({"id": doc.id, "filename": doc.filename, "status": doc.status})
    audit("upload", detail=f"{len(created)} 份{('招标' if kind == 'tender' else '投标')}文件")
    return {"created": created, "queued": len(created)}


def detect_file_type(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext in {".jpg", ".jpeg", ".png"}:
        return "scan"
    if ext == ".txt":
        return "txt"
    return ext.lstrip(".")
